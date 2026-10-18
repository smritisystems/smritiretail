"""
Project      : SMRITI Retail OS
Repository   : SMRITIRetailNX
Organization : AITDL NETWORKS

Founders
- Pushpa Devi Jawahar Mallah (Founder & Chairperson)
- Jawahar Ramkripal Mallah (Founder, CEO & Chief Software Architect)

Websites     : smritibooks.com | erpnbook.com | aitdl.com
Version      : 6.92.3
Created      : 2026-10-19
Modified     : 2026-10-19
Copyright    : © AITDL.com and SMRITIBooks.com. All Rights Reserved.
License      : Proprietary Commercial Software
Classification: Internal
"""

import os
import sys
import openpyxl

EXCEL_PATH = r"F:\Smriti-Clients Data\ItemmasterExported\FullItemmaster.xlsx"

def remove_duplicates():
    print("=" * 80)
    print("Removing Duplicate Columns: HEELS, UPPER MATERIAL, Attribute 3 (A3)")
    print("=" * 80)
    
    wb = openpyxl.load_workbook(EXCEL_PATH)
    ws = wb["Item Master"]
    
    headers = [str(c.value).strip() if c.value is not None else f"COL_{idx}" for idx, c in enumerate(ws[1])]
    print(f"Original total columns: {len(headers)}")
    
    targets_to_delete = ["Attribute 3 (A3)", "UPPER MATERIAL", "HEELS"]
    
    for target in targets_to_delete:
        curr_headers = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
        if target in curr_headers:
            col_idx = curr_headers.index(target) + 1  # 1-indexed
            print(f"  Deleting duplicate column {col_idx}: '{target}'...")
            ws.delete_cols(col_idx, 1)
        else:
            print(f"  Column '{target}' not found (already deleted).")
            
    try:
        wb.save(EXCEL_PATH)
        print(f"\n  --> Saved cleanly to: {EXCEL_PATH}")
    except PermissionError:
        fallback = EXCEL_PATH.replace(".xlsx", "_NoDuplicates.xlsx")
        wb.save(fallback)
        print(f"\n  [NOTE] File was locked, saved to: {fallback}")
        
    # Re-verify
    check_wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    check_ws = check_wb["Item Master"]
    new_headers = [str(c.value).strip() if c.value is not None else "" for c in check_ws[1]]
    print(f"\nNew total columns: {len(new_headers)}")
    print("Clean Column List:")
    for idx, h in enumerate(new_headers):
        print(f"  Col {idx:2}: {h}")

if __name__ == "__main__":
    remove_duplicates()
