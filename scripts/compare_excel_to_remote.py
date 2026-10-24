"""
Project      : SMRITI Retail OS
Repository   : SMRITIRetailNX
Organization : AITDL NETWORKS

Websites     : smritibooks.com | erpnbook.com | aitdl.com
Version      : 6.92.3
Created      : 2026-10-19
Modified     : 2026-10-19
Classification: Internal
"""

import asyncio
import httpx
import openpyxl

EXCEL_PATH = r"F:\Smriti-Clients Data\ItemmasterExported\FullItemmaster.xlsx"
BASE_URL = "https://tattlythreads.smritisys.com"

async def compare_excel_and_remote():
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    ws = wb["Item Master"]
    headers = [str(c.value).strip() if c.value is not None else f"COL_{idx}" for idx, c in enumerate(ws[1])]
    h_map = {name: idx for idx, name in enumerate(headers)}
    
    excel_rows = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not any(r): continue
        sku = str(r[h_map["Stock No / SKU"]]).strip().upper() if r[h_map["Stock No / SKU"]] else ""
        bc = str(r[h_map["Barcode"]]).strip() if r[h_map["Barcode"]] else ""
        entry = {h: r[idx] for h, idx in h_map.items()}
        if sku:
            excel_rows[sku] = entry
        if bc:
            excel_rows[bc] = entry
        
    print(f"Excel rows indexed: {len(excel_rows)}")
    
    async with httpx.AsyncClient(timeout=60.0) as client:
        login_res = await client.post(f"{BASE_URL}/api/v1/auth/login", json={"username": "admin", "password": "Admin@123"})
        token = login_res.json()["access_token"]
        headers = {"Authorization": f"Bearer {token}"}
        
        all_prods = []
        page = 1
        while True:
            resp = await client.get(f"{BASE_URL}/api/v1/inventory/?page={page}&page_size=100", headers=headers)
            items = resp.json().get("items", [])
            if not items: break
            all_prods.extend(items)
            if not resp.json().get("has_next") or len(all_prods) >= resp.json().get("total", 0): break
            page += 1
            
        print(f"Remote products loaded: {len(all_prods)}")
        
        diff_count = 0
        diff_samples = []
        for p in all_prods:
            sku = (p.get("code") or p.get("sku") or "").upper()
            bc = p.get("barcode") or ""
            ex_row = excel_rows.get(sku) or excel_rows.get(bc)
            if not ex_row:
                continue
                
            attrs = p.get("attributes") or {}
            
            # Compare key dimensions
            checks = [
                ("Vendor Code", p.get("vendor_code") or attrs.get("vendor_code"), ex_row.get("Vendor Code")),
                ("Department", attrs.get("department"), ex_row.get("Department")),
                ("Sub-Category", attrs.get("sub_category"), ex_row.get("Sub-Category")),
                ("Heel Type", attrs.get("heel_type"), ex_row.get("Heel Type")),
                ("Upper Material", attrs.get("upper_material"), ex_row.get("Upper Material")),
                ("Outsole Material", attrs.get("outsole_material") or attrs.get("outsole"), ex_row.get("Outsole Material")),
                ("Purchase Class", attrs.get("purchase_class"), ex_row.get("Purchase Class")),
                ("Image Link", attrs.get("image_link"), ex_row.get("Image Name")),
            ]
            
            p_diffs = []
            for field, r_val, e_val in checks:
                r_s = str(r_val).strip().upper() if r_val is not None else ""
                e_s = str(e_val).strip().upper() if e_val is not None else ""
                if r_s != e_s and e_s != "":
                    p_diffs.append((field, r_s, e_s))
                    
            if p_diffs:
                diff_count += 1
                if len(diff_samples) < 10:
                    diff_samples.append((sku, p_diffs))
                    
        print(f"Products with differences between Excel and Remote: {diff_count} / {len(all_prods)}")
        for s, diffs in diff_samples:
            print(f"  Product {s}:")
            for f, r_val, e_val in diffs:
                print(f"    - {f}: Remote=\"{r_val}\" vs Excel=\"{e_val}\"")

if __name__ == "__main__":
    asyncio.run(compare_excel_and_remote())
