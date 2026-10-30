"""
Project      : SMRITI Retail OS
Repository   : SMRITIRetailNX
Organization : AITDL NETWORKS

Founders
- Pushpa Devi Jawahar Mallah (Founder & Chairperson)
- Jawahar Ramkripal Mallah (Founder, CEO & Chief Software Architect)

Websites     : smritibooks.com | erpnbook.com | aitdl.com
Version      : 6.92.3
Created      : 2026-10-19
Modified     : 2026-10-19
Copyright    : © AITDL.com and SMRITIBooks.com. All Rights Reserved.
License      : Proprietary Commercial Software
Classification: Internal
"""

import os
import sys
import json
import asyncio
import uuid
import httpx
import psycopg2
import openpyxl

sys.stdout.reconfigure(encoding="utf-8")

DIR_PATH = r"F:\Smriti-Clients Data\Tattly Threads\ItemMaster"
BASE_URL = "https://tattlythreads.smritisys.com"

def clean_val(v):
    if v is None:
        return ""
    s = str(v).strip()
    return "" if s.lower() in ("none", "null") else s

def clean_num(v, default=0.0):
    if v is None:
        return default
    try:
        return float(str(v).replace(",", "").strip())
    except Exception:
        return default

def load_all_excel_itemmasters():
    print("================================================================================")
    print("STEP 1: Compiling Full Item Master Dataset from Excel Workbooks")
    print("================================================================================")
    records = {}
    styles_meta = {}

    # 1. 1136 file
    fp1 = os.path.join(DIR_PATH, "TATTLY_THREADS_ALL_1136_BARCODES_WITH_STYLES.xlsx")
    if os.path.exists(fp1):
        print(f"  Reading {os.path.basename(fp1)}...")
        wb = openpyxl.load_workbook(fp1, data_only=True)
        ws = wb["Barcodes & Styles (1136)"]
        for row in ws.iter_rows(min_row=2, values_only=True):
            bc = clean_val(row[1])
            if not bc:
                continue
            st = clean_val(row[2])
            sku = clean_val(row[5])
            col = clean_val(row[6]).upper()
            sz = clean_val(row[7])
            br = clean_val(row[8]) or "Tattly Threads"
            cat = clean_val(row[9]).upper() or "SANDAL"
            gen = clean_val(row[10]).upper() or "LADIES"
            dep = clean_val(row[11]).upper() or "LADIES FTW"
            mrp = clean_num(row[12], 1299.0)
            cost = clean_num(row[13], 350.0)
            gst = clean_num(row[15], 5.0)
            hsn = clean_val(row[16]) or "64041990"
            heel = clean_val(row[17]).upper() or "FLAT"
            upper = clean_val(row[18]).upper() or "SYNTHETIC"
            sole = clean_val(row[19]).upper() or "TPR"

            entry = {
                "barcode": bc, "style": st, "sku": sku, "color": col, "size": sz, "brand": br,
                "category": cat, "merchandise_category": cat, "product_type": cat, "gender": gen,
                "department": dep, "mrp": mrp, "cost_price": cost, "buying_price": cost,
                "gst_percentage": gst, "hsn_code": hsn, "heel_type": heel, "upper_material": upper,
                "outsole": sole, "outsole_material": sole, "size_group": "FOOTWEAR_EU",
                "tax_category": "GOODS", "uom": "Pair", "collection_type": "BASIC",
                "purchase_class": "SIS",
            }
            records[bc] = entry
            if st:
                styles_meta[st.upper()] = entry

    # 2. TATTLY THREADS ITEM MASTER 2026 - Copy.xlsx
    fp2 = os.path.join(DIR_PATH, "TATTLY THREADS ITEM MASTER 2026 - Copy.xlsx")
    if os.path.exists(fp2):
        print(f"  Reading {os.path.basename(fp2)}...")
        wb = openpyxl.load_workbook(fp2, data_only=True)
        for s in ["NEW", "Sheet2"]:
            if s in wb.sheetnames:
                ws = wb[s]
                h = [str(c.value).strip() if c.value else "" for c in ws[1]]
                h_map = {name: idx for idx, name in enumerate(h) if name}
                for row in ws.iter_rows(min_row=2, values_only=True):
                    bc = clean_val(row[h_map.get("BARCODE NO", 0)])
                    if not bc:
                        continue
                    st = clean_val(row[h_map.get("PRODUCT STYLE CODE", 1)])
                    desc = clean_val(row[h_map.get("ITEM DESCRIPTION", 2)]).upper() or "BASIC"
                    br = clean_val(row[h_map.get("BRAND NAME", 3)]) or "Tattly Threads"
                    col = clean_val(row[h_map.get("COLOR", 4)]).upper()
                    sz = clean_val(row[h_map.get("SIZE", 5)])
                    mrp = clean_num(row[h_map.get("PLANNED MRP", 7)], 1299.0)
                    cost = clean_num(row[h_map.get("COST PRICE", 8)], 350.0)
                    gst = clean_num(row[h_map.get("PRODUCT TAX", 9)], 5.0)
                    hsn = clean_val(row[h_map.get("HSN CODE", 10)]) or "64041990"
                    gen = clean_val(row[h_map.get("GENDER", 11)]).upper() or "LADIES"
                    vc = clean_val(row[h_map.get("VENDOR CODE", 12)])
                    pc = clean_val(row[h_map.get("PURCHASE CLASS", 13)]).upper() or "SIS"
                    dep = clean_val(row[h_map.get("DEPARTMENT", 14)]).upper() or "LADIES FTW"
                    cat = clean_val(row[h_map.get("MERCHANDISE CATEGORY", 15)]).upper() or "CHAPPAL"
                    subcat = clean_val(row[h_map.get("Sub category", 16)]).upper() or "STANDARD"
                    heel = clean_val(row[h_map.get("HEELS", 17)]).upper() or "FLAT"
                    upper = clean_val(row[h_map.get("UPPER MATERIAL", 18)]).upper() or "SYNTHETIC"
                    sole = clean_val(row[h_map.get("OUTSOLE", 19)]).upper() or "TPR"
                    img = clean_val(row[h_map.get("IMAGE LINK", 20)])

                    entry = {
                        "barcode": bc, "style": st, "sku": f"{st}-{col}-{sz}", "item_description": desc,
                        "collection_type": desc, "brand": br, "color": col, "size": sz, "mrp": mrp,
                        "cost_price": cost, "buying_price": cost, "gst_percentage": gst, "hsn_code": hsn,
                        "gender": gen, "vendor_code": vc, "purchase_class": pc, "department": dep,
                        "merchandise_category": cat, "category": cat, "product_type": cat,
                        "sub_category": subcat, "design_attribute": subcat, "heel_type": heel,
                        "upper_material": upper, "outsole": sole, "outsole_material": sole,
                        "image_link": img, "size_group": "FOOTWEAR_EU", "tax_category": "GOODS", "uom": "Pair",
                    }
                    if bc not in records:
                        records[bc] = entry
                    else:
                        for k, v in entry.items():
                            if v and not records[bc].get(k):
                                records[bc][k] = v
                    if st:
                        st_u = st.upper()
                        if st_u not in styles_meta:
                            styles_meta[st_u] = entry
                        else:
                            for k, v in entry.items():
                                if v and not styles_meta[st_u].get(k):
                                    styles_meta[st_u][k] = v

    # 3. OLD TO NEW.xlsx
    for fn in ["OLD TO NEW.xlsx", "OLD TO NEW (1).xlsx"]:
        fp3 = os.path.join(DIR_PATH, fn)
        if os.path.exists(fp3):
            print(f"  Reading {fn}...")
            wb = openpyxl.load_workbook(fp3, data_only=True)
            for s in wb.sheetnames:
                ws = wb[s]
                h = [str(c.value).strip() if c.value else "" for c in ws[1]]
                h_map = {name: idx for idx, name in enumerate(h) if name}
                if "BARCODE NO" in h_map:
                    for row in ws.iter_rows(min_row=2, values_only=True):
                        bc = clean_val(row[h_map["BARCODE NO"]])
                        if not bc:
                            continue
                        st = clean_val(row[h_map.get("PRODUCT STYLE CODE", 1)])
                        col = clean_val(row[h_map.get("COLOR", 4)]).upper()
                        sz = clean_val(row[h_map.get("SIZE", 5)])
                        mrp = clean_num(row[h_map.get("PLANNED MRP", 6)], 1299.0)
                        cost = clean_num(row[h_map.get("COST PRICE", 7)], 350.0)
                        subcat = clean_val(row[h_map.get("Sub category", 15)]).upper() or "STANDARD"
                        heel = clean_val(row[h_map.get("HEELS", 16)]).upper() or "FLAT"
                        upper = clean_val(row[h_map.get("UPPER MATERIAL", 17)]).upper() or "SYNTHETIC"
                        sole = clean_val(row[h_map.get("OUTSOLE", 18)]).upper() or "TPR"
                        img = clean_val(row[h_map.get("IMAGE LINK", 19)])

                        entry = {
                            "barcode": bc, "style": st, "sku": f"{st}-{col}-{sz}", "brand": "Tattly Threads",
                            "color": col, "size": sz, "mrp": mrp, "cost_price": cost, "buying_price": cost,
                            "gst_percentage": 5.0, "hsn_code": "64041990", "gender": "LADIES",
                            "department": "LADIES FTW", "merchandise_category": "CHAPPAL", "category": "CHAPPAL",
                            "product_type": "CHAPPAL", "sub_category": subcat, "design_attribute": subcat,
                            "heel_type": heel, "upper_material": upper, "outsole": sole,
                            "outsole_material": sole, "image_link": img, "size_group": "FOOTWEAR_EU",
                            "tax_category": "GOODS", "uom": "Pair", "collection_type": "BASIC",
                            "purchase_class": "SIS",
                        }
                        if bc not in records:
                            records[bc] = entry
                        else:
                            for k, v in entry.items():
                                if v and not records[bc].get(k):
                                    records[bc][k] = v

    print(f"  --> Total compiled unique barcodes : {len(records)}")
    print(f"  --> Total compiled unique styles   : {len(styles_meta)}")
    return records, styles_meta

def match_product_record(prod, records, styles_meta):
    bc = clean_val(prod.get("barcode"))
    code = clean_val(prod.get("code")).upper()
    sku = clean_val(prod.get("sku")).upper()

    if bc and bc in records:
        return records[bc]
    for r in records.values():
        if code and (r.get("sku", "").upper() == code or r.get("barcode") == code):
            return r
        if sku and r.get("sku", "").upper() == sku:
            return r

    # Match by style prefix
    parts = code.split("-")
    if len(parts) >= 2:
        two_part = f"{parts[0]}-{parts[1]}"
        if two_part in styles_meta:
            return styles_meta[two_part]
        if len(parts) >= 3:
            three_part = f"{parts[0]}-{parts[1]}-{parts[2]}"
            if three_part in styles_meta:
                return styles_meta[three_part]
    if parts and parts[0] in styles_meta:
        return styles_meta[parts[0]]

    return None

async def execute_live_update(records, styles_meta):
    print("\n================================================================================")
    print("STEP 2: Synchronizing with Remote Production Server https://tattlythreads.smritisys.com")
    print("================================================================================")

    async with httpx.AsyncClient(timeout=120.0) as client:
        login_res = await client.post(
            f"{BASE_URL}/api/v1/auth/login",
            json={"username": "admin", "password": "Admin@123"},
        )
        if login_res.status_code != 200:
            print(f"[FATAL] Login failed: {login_res.status_code} {login_res.text}")
            return
        token = login_res.json()["access_token"]
        headers = {"Authorization": f"Bearer {token}", "Content-Type": "application/json"}
        print("  Successfully authenticated as Admin.")

        # 1. Fetch all existing products on remote
        print("  Scanning all existing remote products...")
        all_remote = []
        page = 1
        while True:
            resp = await client.get(
                f"{BASE_URL}/api/v1/inventory/?page={page}&page_size=100",
                headers=headers,
            )
            data = resp.json()
            items = data.get("items", [])
            if not items:
                break
            all_remote.extend(items)
            if not data.get("has_next") or len(all_remote) >= data.get("total", 0):
                break
            page += 1

        print(f"  --> Found {len(all_remote)} existing products on remote server.")
        remote_barcodes = {clean_val(p.get("barcode")) for p in all_remote if p.get("barcode")}

        # Sync master lookups (styles and vendor codes)
        style_res = await client.get(f"{BASE_URL}/api/v1/lookup/style_article/values", headers=headers)
        reg_styles = {x.get("code") for x in style_res.json() if x.get("code")} if style_res.status_code == 200 else set()
        needed_styles = set(styles_meta.keys())
        for p in all_remote:
            c = clean_val(p.get("code")).upper()
            sk = clean_val(p.get("sku")).upper()
            st = clean_val(p.get("style_code")).upper()
            if st:
                needed_styles.add(st)
            parts = (c or sk).split("-")
            if len(parts) >= 3 and len(parts[2]) == 1:
                needed_styles.add(f"{parts[0]}-{parts[1]}-{parts[2]}")
            elif len(parts) >= 2:
                needed_styles.add(f"{parts[0]}-{parts[1]}")
            elif parts:
                needed_styles.add(parts[0])

        missing_styles = sorted([s for s in needed_styles if s and s not in reg_styles])
        if missing_styles:
            print(f"  Registering {len(missing_styles)} missing styles in remote style_article lookup...")
            for st in missing_styles:
                res = await client.post(
                    f"{BASE_URL}/api/v1/lookup/style_article/values",
                    headers=headers,
                    json={"code": st, "name": st, "active": True, "data": {}},
                )
                if res.status_code in (200, 201):
                    reg_styles.add(st)

        # 2. Update all existing products on remote with complete footwear attributes
        print("\n  [Phase A] Updating attributes on all existing remote products...")
        mutated_count = 0
        skipped_count = 0
        semaphore = asyncio.Semaphore(15)

        async def update_single_product(prod):
            nonlocal mutated_count, skipped_count
            meta = match_product_record(prod, records, styles_meta)
            if not meta:
                skipped_count += 1
                return

            curr_attrs = dict(prod.get("attributes") or {})
            changed = False

            attr_keys = [
                "heel_type", "upper_material", "outsole", "outsole_material",
                "merchandise_category", "sub_category", "design_attribute",
                "collection_type", "department", "gender", "purchase_class",
                "image_link", "vendor_code"
            ]

            for k in attr_keys:
                val = meta.get(k)
                if val and curr_attrs.get(k) != val:
                    curr_attrs[k] = val
                    changed = True

            put_body = {"attributes": curr_attrs}
            st = meta.get("style") or clean_val(prod.get("style_code"))
            if st and st in reg_styles:
                put_body["style_code"] = st

            # Update core fields if empty or differ
            if meta.get("brand") and not prod.get("brand"):
                put_body["brand"] = meta["brand"]
                changed = True
            if meta.get("vendor_code") and clean_val(prod.get("vendor_code")) != meta["vendor_code"]:
                put_body["vendor_code"] = meta["vendor_code"]
                changed = True
            if meta.get("hsn_code") and clean_val(prod.get("hsn_code")) != meta["hsn_code"]:
                put_body["hsn_code"] = meta["hsn_code"]
                changed = True

            if not changed:
                skipped_count += 1
                return

            pid = prod.get("id")
            async with semaphore:
                try:
                    put_res = await client.put(
                        f"{BASE_URL}/api/v1/inventory/{pid}",
                        headers=headers,
                        json=put_body,
                    )
                    if put_res.status_code == 200:
                        mutated_count += 1
                    else:
                        print(f"    Warning: PUT {pid} failed: {put_res.status_code}")
                except Exception as err:
                    print(f"    Error PUT {pid}: {err}")

        tasks = [update_single_product(p) for p in all_remote]
        await asyncio.gather(*tasks)

        print(f"  --> Updated {mutated_count} existing remote products (Skipped/Identical: {skipped_count}).")

        # 3. Identify and import any barcodes in Excel that do not exist yet on remote
        missing_barcodes = [bc for bc, r in records.items() if bc and bc not in remote_barcodes]
        print(f"\n  [Phase B] Identifying new Excel barcodes not yet in remote inventory...")
        print(f"  --> Found {len(missing_barcodes)} new barcodes to import via Universal Import.")

        if missing_barcodes:
            # Build Universal Import payload
            import_rows = []
            for idx, bc in enumerate(missing_barcodes, 1):
                r = records[bc]
                style = r.get("style") or "CH-01"
                color = r.get("color") or "BLACK"
                size = r.get("size") or "36"
                sku = r.get("sku") or f"{style}-{color}-{sz}"
                cat = r.get("category") or "CHAPPAL"
                if cat.upper() == "SHOES":
                    cat = "SHOE"

                row_dict = {
                    "rowNumber": idx,
                    "barcode": bc,
                    "sku": sku,
                    "style_code": style,
                    "product_style_code": style,
                    "product_name": f"Tattly Threads {style} {color} {size}",
                    "brand": r.get("brand") or "Tattly Threads",
                    "color": color,
                    "size": str(size),
                    "size_group": "FOOTWEAR_EU",
                    "tax_category": "GOODS",
                    "mrp": float(r.get("mrp") or 1299.0),
                    "selling_price": float(r.get("mrp") or 1299.0),
                    "cost_price": float(r.get("cost_price") or 350.0),
                    "buying_price": float(r.get("cost_price") or 350.0),
                    "gst_percentage": float(r.get("gst_percentage") or 5.0),
                    "hsn_code": r.get("hsn_code") or "64041990",
                    "gender": r.get("gender") or "LADIES",
                    "department": r.get("department") or "LADIES FTW",
                    "category": cat,
                    "merchandise_category": cat,
                    "product_type": cat,
                    "collection_type": r.get("collection_type") or "BASIC",
                    "subCategory": r.get("sub_category") or "STANDARD",
                    "design_attribute": r.get("design_attribute") or "STANDARD",
                    "heel_type": r.get("heel_type") or "FLAT",
                    "upper_material": r.get("upper_material") or "SYNTHETIC",
                    "outsole_material": r.get("outsole_material") or "TPR",
                    "uom": "Pair",
                }
                import_rows.append(row_dict)

            # Preview
            print("  Executing Universal Import Preview for new items...")
            preview_res = await client.post(
                f"{BASE_URL}/api/v1/universal-import/preview",
                json={"target": "ITEM_MASTER", "rows": import_rows},
                headers=headers,
            )
            if preview_res.status_code == 200:
                p_data = preview_res.json()
                summary = p_data.get("summary") or {}
                print(f"    Preview Status: {summary.get('status')}")
                print(f"    Valid Rows: {summary.get('valid_rows')} / {summary.get('total_rows')}")
                print(f"    Blocking Errors: {summary.get('blocking_errors')}")

                valid_rows = [r.get("raw_row") for r in p_data.get("rows", []) if r.get("status") != "INVALID" and r.get("raw_row")]
                if not valid_rows:
                    valid_rows = import_rows

                # Commit in batches of 100
                batch_size = 100
                total_created = 0
                for i in range(0, len(valid_rows), batch_size):
                    batch = valid_rows[i : i + batch_size]
                    commit_payload = {
                        "target": "ITEM_MASTER",
                        "rows": batch,
                        "strategy": "ALL_ELIGIBLE",
                        "import_strategy": "ALL_ELIGIBLE",
                        "idempotency_key": f"import_{uuid.uuid4().hex[:10]}",
                        "dry_run": False,
                    }
                    commit_res = await client.post(
                        f"{BASE_URL}/api/v1/universal-import/commit",
                        json=commit_payload,
                        headers=headers,
                    )
                    if commit_res.status_code == 200:
                        c_data = commit_res.json()
                        created_in_batch = c_data.get("created", len(batch))
                        total_created += created_in_batch
                        print(f"    Batch committed successfully: {created_in_batch} items created.")
                    else:
                        print(f"    Warning: Commit batch failed ({commit_res.status_code}): {commit_res.text[:300]}")

                print(f"  --> Successfully created {total_created} new products on remote server!")
            else:
                print(f"    Preview failed ({preview_res.status_code}): {preview_res.text[:300]}")

        # Final verification
        print("\n================================================================================")
        print("FINAL VERIFICATION: Checking Random Samples on Live Server")
        print("================================================================================")
        test_samples = ["8904551000002", "8904551003089", "8904551003577", "8904551001962", "8904551002037"]
        for bc in test_samples:
            check_res = await client.get(f"{BASE_URL}/api/v1/inventory/?search={bc}", headers=headers)
            if check_res.status_code == 200:
                items = check_res.json().get("items", [])
                if items:
                    p = items[0]
                    attrs = p.get("attributes", {})
                    print(f"  Barcode {bc} ({p.get('code')}) => Heel: {attrs.get('heel_type')} | Upper: {attrs.get('upper_material')} | Outsole: {attrs.get('outsole_material') or attrs.get('outsole')} | Cat: {attrs.get('merchandise_category')} | SubCat: {attrs.get('sub_category')}")

def update_local_db(records, styles_meta):
    print("\n================================================================================")
    print("STEP 3: Synchronizing Local PostgreSQL Database (localhost:2781/smritisys)")
    print("================================================================================")
    try:
        conn = psycopg2.connect("postgresql://postgres:postgres@localhost:2781/smritisys")
        cur = conn.cursor()
        cur.execute("SELECT id, code, barcode, sku, style_code, gst_percentage, hsn_code, vendor_code, category, attributes FROM products;")
        rows = cur.fetchall()
        print(f"  Fetched {len(rows)} local products.")
        updated = 0
        for pid, code, barcode, sku, style_code, gst_pct, hsn, vc, cat, attrs in rows:
            prod = {"id": pid, "code": code, "barcode": barcode, "sku": sku, "style_code": style_code}
            entry = match_product_record(prod, records, styles_meta)
            if entry:
                curr_attrs = dict(attrs or {})
                changed = False
                for k in ["heel_type", "upper_material", "outsole", "outsole_material", "merchandise_category", "sub_category", "design_attribute", "collection_type", "department", "gender", "purchase_class", "image_link", "vendor_code"]:
                    v = entry.get(k)
                    if v and curr_attrs.get(k) != v:
                        curr_attrs[k] = v
                        changed = True

                update_cols = {}
                if gst_pct != 5.0:
                    update_cols["gst_percentage"] = 5.0
                if entry.get("hsn_code") and clean_val(hsn) != entry["hsn_code"]:
                    update_cols["hsn_code"] = entry["hsn_code"]
                if entry.get("vendor_code") and clean_val(vc) != entry["vendor_code"]:
                    update_cols["vendor_code"] = entry["vendor_code"]
                if entry.get("merchandise_category") and clean_val(cat).upper() != entry["merchandise_category"].upper():
                    update_cols["category"] = entry["merchandise_category"]
                if entry.get("style") and (not style_code or style_code != entry["style"]):
                    update_cols["style_code"] = entry["style"]

                if changed or update_cols:
                    set_clauses = ["attributes = %s"]
                    params = [json.dumps(curr_attrs)]
                    for col_name, col_val in update_cols.items():
                        set_clauses.append(f"{col_name} = %s")
                        params.append(col_val)
                    params.append(pid)

                    cur.execute(
                        f"UPDATE products SET {', '.join(set_clauses)} WHERE id = %s;",
                        tuple(params),
                    )
                    updated += 1
        conn.commit()
        conn.close()
        print(f"  --> Local PostgreSQL updated successfully! ({updated} products modified)")
    except Exception as e:
        print(f"  Local DB update exception: {e}")

def main():
    records, styles_meta = load_all_excel_itemmasters()
    update_local_db(records, styles_meta)
    asyncio.run(execute_live_update(records, styles_meta))

if __name__ == "__main__":
    main()
