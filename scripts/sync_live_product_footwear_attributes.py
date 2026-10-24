"""
Project      : SMRITI Retail OS
Repository   : SMRITIRetailNX
Organization : AITDL NETWORKS

Founders
- Pushpa Devi Jawahar Mallah (Founder & Chairperson)
- Jawahar Ramkripal Mallah (Founder, CEO & Chief Software Architect)

Websites: aitdl.com | erpnbook.com | smritibooks.com

Version    : 3.35.0
Created    : 2026-10-19
Modified   : 2026-10-19
Copyright  : © AITDL.com and SMRITIBooks.com. All Rights Reserved.
License    : Proprietary Commercial Software
Classification: Internal
"""

import sys
import json
import asyncio
import httpx
import psycopg2
import openpyxl

sys.stdout.reconfigure(encoding="utf-8")

def build_attribute_catalogs():
    print("[1/4] Loading Excel workbooks...")
    wb1 = openpyxl.load_workbook(r"F:\Smriti-Clients Data\Tattly Threads\ItemMaster\TATTLY_THREADS_ALL_1136_BARCODES_WITH_STYLES.xlsx", data_only=True)
    ws1 = wb1["Barcodes & Styles (1136)"]

    by_barcode = {}
    by_sku = {}
    by_style = {}

    for row in ws1.iter_rows(min_row=2, values_only=True):
        barcode = str(row[1]).strip() if row[1] else ""
        style = str(row[2]).strip() if row[2] else ""
        sku = str(row[5]).strip() if row[5] else ""
        color = str(row[6]).strip() if row[6] else ""
        size = str(row[7]).strip() if row[7] else ""
        brand = str(row[8]).strip() if row[8] else "Tattly Threads"
        category = str(row[9]).strip() if row[9] else ""
        gender = str(row[10]).strip() if row[10] else "LADIES"
        dept = str(row[11]).strip() if row[11] else "LADIES FTW"
        hsn = str(row[16]).strip() if row[16] else ""
        heel = str(row[17]).strip() if row[17] else ""
        upper = str(row[18]).strip() if row[18] else ""
        sole = str(row[19]).strip() if row[19] else ""

        entry = {
            "heel_type": heel,
            "upper_material": upper,
            "outsole": sole,
            "outsole_material": sole,
            "category": category,
            "merchandise_category": category,
            "gender": gender,
            "department": dept,
            "brand": brand,
            "hsn_code": hsn,
            "color": color,
            "size": size,
            "purchase_class": "SIS",
        }

        if barcode:
            by_barcode[barcode] = entry
        if sku:
            by_sku[sku.upper()] = entry
        if style:
            clean_s = style.upper()
            if clean_s not in by_style:
                by_style[clean_s] = entry
            # Also register vendor sub-prefixes (A, B, C, J, etc.)
            for v in ["A", "B", "C", "J", "N", "R"]:
                by_style[f"{clean_s}-{v}"] = entry

    wb2 = openpyxl.load_workbook(r"F:\Smriti-Clients Data\Tattly Threads\ItemMaster\TATTLY THREADS ITEM MASTER 2026 - Copy.xlsx", data_only=True)
    ws2 = wb2["NEW"]
    for row in ws2.iter_rows(min_row=2, values_only=True):
        barcode = str(row[0]).strip() if row[0] else ""
        style = str(row[1]).strip() if row[1] else ""
        subcat = str(row[16]).strip() if row[16] else ""
        heel = str(row[17]).strip() if row[17] else ""
        upper = str(row[18]).strip() if row[18] else ""
        sole = str(row[19]).strip() if row[19] else ""
        cat = str(row[15]).strip() if row[15] else ""

        if barcode and barcode in by_barcode:
            if subcat:
                by_barcode[barcode]["sub_category"] = subcat
                by_barcode[barcode]["design_attribute"] = subcat

        if style:
            s_upper = style.upper()
            s_base = style.split("-A")[0].strip().upper()
            for k in [s_upper, s_base]:
                if k in by_style:
                    if subcat:
                        by_style[k]["sub_category"] = subcat
                        by_style[k]["design_attribute"] = subcat
                    if heel and not by_style[k].get("heel_type"):
                        by_style[k]["heel_type"] = heel
                    if upper and not by_style[k].get("upper_material"):
                        by_style[k]["upper_material"] = upper
                    if sole and not by_style[k].get("outsole"):
                        by_style[k]["outsole"] = sole
                        by_style[k]["outsole_material"] = sole

    print(f"  Loaded {len(by_barcode)} barcodes, {len(by_sku)} SKUs, {len(by_style)} styles.")
    return by_barcode, by_sku, by_style

def resolve_meta(prod, by_barcode, by_sku, by_style):
    barcode = str(prod.get("barcode") or "").strip()
    code = str(prod.get("code") or "").strip().upper()
    sku = str(prod.get("sku") or "").strip().upper()

    if barcode and barcode in by_barcode:
        return by_barcode[barcode]
    if code and code in by_sku:
        return by_sku[code]
    if sku and sku in by_sku:
        return by_sku[sku]

    # Try style prefix
    parts = code.split("-")
    if len(parts) >= 2:
        two_part = f"{parts[0]}-{parts[1]}"
        if two_part in by_style:
            return by_style[two_part]
        if len(parts) >= 3:
            three_part = f"{parts[0]}-{parts[1]}-{parts[2]}"
            if three_part in by_style:
                return by_style[three_part]

    for s_name, meta in by_style.items():
        if s_name in code or code in s_name:
            return meta

    return None

def update_local_postgres(by_barcode, by_sku, by_style):
    print("[2/4] Updating Local PostgreSQL...")
    try:
        conn = psycopg2.connect("postgresql://postgres:postgres@localhost:2781/smritisys")
        cur = conn.cursor()
        cur.execute("SELECT id, code, barcode, attributes FROM products;")
        rows = cur.fetchall()
        print(f"  Fetched {len(rows)} local products.")
        updated_count = 0

        for pid, code, barcode, attrs in rows:
            prod = {"id": pid, "code": code, "barcode": barcode}
            meta = resolve_meta(prod, by_barcode, by_sku, by_style)
            if meta:
                curr_attrs = dict(attrs or {})
                changed = False
                for k, v in meta.items():
                    if v and curr_attrs.get(k) != v:
                        curr_attrs[k] = v
                        changed = True
                if changed:
                    cur.execute(
                        "UPDATE products SET attributes = %s WHERE id = %s;",
                        (json.dumps(curr_attrs), pid),
                    )
                    updated_count += 1

        conn.commit()
        conn.close()
        print(f"  Local PostgreSQL updated successfully! ({updated_count} products modified)")
    except Exception as e:
        print(f"  Local PostgreSQL update error: {e}")

async def update_remote_server(by_barcode, by_sku, by_style):
    print("[3/4] Authenticating with Remote Server https://tattlythreads.smritisys.com...")
    base_url = "https://tattlythreads.smritisys.com"
    async with httpx.AsyncClient(timeout=60.0) as client:
        login_res = await client.post(
            f"{base_url}/api/v1/auth/login",
            json={"username": "admin", "password": "Admin@123"},
        )
        if login_res.status_code != 200:
            print(f"  Failed remote login: {login_res.status_code} {login_res.text}")
            return
        token = login_res.json().get("access_token")
        headers = {"Authorization": f"Bearer {token}"}
        print("  Remote login successful. Scanning all remote products...")

        page = 1
        page_size = 100
        all_prods = []
        while True:
            resp = await client.get(
                f"{base_url}/api/v1/inventory/?page={page}&page_size={page_size}",
                headers=headers,
            )
            if resp.status_code != 200:
                print(f"  Failed fetching page {page}: {resp.status_code}")
                break
            data = resp.json()
            items = data.get("items", [])
            if not items:
                break
            all_prods.extend(items)
            print(f"    Page {page}: fetched {len(items)} products (total: {len(all_prods)}/{data.get('total')})")
            if not data.get("has_next") or len(all_prods) >= data.get("total", 0):
                break
            page += 1

        print(f"[4/4] Updating {len(all_prods)} remote products with complete footwear attributes...")
        mutated_count = 0
        skipped_count = 0
        semaphore = asyncio.Semaphore(10)

        async def update_single_product(prod):
            nonlocal mutated_count, skipped_count
            meta = resolve_meta(prod, by_barcode, by_sku, by_style)
            if not meta:
                skipped_count += 1
                return

            curr_attrs = dict(prod.get("attributes") or {})
            changed = False
            for k in ["heel_type", "upper_material", "outsole", "outsole_material", "merchandise_category", "sub_category", "department", "gender", "purchase_class"]:
                val = meta.get(k)
                if val and curr_attrs.get(k) != val:
                    curr_attrs[k] = val
                    changed = True

            if not changed:
                skipped_count += 1
                return

            pid = prod.get("id")
            async with semaphore:
                try:
                    put_res = await client.put(
                        f"{base_url}/api/v1/inventory/{pid}",
                        headers=headers,
                        json={"attributes": curr_attrs},
                    )
                    if put_res.status_code == 200:
                        mutated_count += 1
                    else:
                        print(f"    Warning: PUT {pid} failed: {put_res.status_code}")
                except Exception as err:
                    print(f"    Error PUT {pid}: {err}")

        tasks = [update_single_product(p) for p in all_prods]
        await asyncio.gather(*tasks)

        print(f"\n=======================================================")
        print(f"REMOTE SYNC COMPLETE:")
        print(f"  Total scanned : {len(all_prods)}")
        print(f"  Updated/Mutated: {mutated_count}")
        print(f"  Unchanged/Skip: {skipped_count}")
        print(f"=======================================================")

        # Verify PO-J-22 on remote
        print("\nVerifying PO-J-22 on Remote Server...")
        po_search = await client.get(
            f"{base_url}/api/v1/purchase/orders?search=PO-J-22",
            headers=headers,
        )
        if po_search.status_code == 200:
            for po in po_search.json():
                if "PO-J-22" in str(po.get("order_no")):
                    po_full = (await client.get(
                        f"{base_url}/api/v1/purchase/orders/{po.get('id')}",
                        headers=headers,
                    )).json()
                    print(f"PO {po_full.get('order_no')} ({po_full.get('id')}) status: {po_full.get('status')}")
                    for it in (po_full.get("items") or [])[:4]:
                        print(f"  Item: {it.get('code')} => Heel: {it.get('heel_type')} | Upper: {it.get('upper_material')} | Outsole: {it.get('outsole_material')}")

def main():
    by_barcode, by_sku, by_style = build_attribute_catalogs()
    update_local_postgres(by_barcode, by_sku, by_style)
    asyncio.run(update_remote_server(by_barcode, by_sku, by_style))

if __name__ == "__main__":
    main()
