import openpyxl
import os
import sys
from collections import defaultdict, Counter

sys.stdout.reconfigure(encoding='utf-8')

folder = r"F:\Smriti-Clients Data\Tattly Threads\ItemMaster"
fields = [
    'PRODUCT TAX', 'HSN CODE', 'GENDER', 'VENDOR CODE', 
    'PURCHASE CLASS', 'DEPARTMENT', 'MERCHANDISE CATEGORY', 
    'Sub category', 'HEELS', 'UPPER MATERIAL', 'OUTSOLE', 'IMAGE LINK'
]

for fname in ['TATTLY THREADS ITEM MASTER 2026 - Copy.xlsx', 'OLD TO NEW.xlsx', 'OLD TO NEW (1).xlsx']:
    fpath = os.path.join(folder, fname)
    if not os.path.exists(fpath):
        continue
    wb = openpyxl.load_workbook(fpath, data_only=True)
    print('========================================================')
    print('FILE:', fname)
    for sname in wb.sheetnames:
        ws = wb[sname]
        h = [str(c.value).strip() if c.value else '' for c in ws[1]]
        h_map = {name: idx for idx, name in enumerate(h) if name}
        if 'PRODUCT STYLE CODE' not in h_map:
            continue
        
        style_col = h_map['PRODUCT STYLE CODE']
        by_style = defaultdict(lambda: defaultdict(list))
        for r_idx, r in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            if not any(r): continue
            st = str(r[style_col]).strip() if r[style_col] else ''
            if not st: continue
            for f in fields:
                if f in h_map:
                    v = r[h_map[f]]
                    val_str = str(v).strip() if v is not None and str(v).strip() != '' and str(v).strip().lower() != 'none' else None
                    by_style[st][f].append((r_idx, val_str))
        
        diffs_found = []
        blanks_found = []
        for st, f_dict in by_style.items():
            for f in fields:
                if f not in f_dict: continue
                pairs = f_dict[f]
                non_blanks = [v for _, v in pairs if v is not None]
                blanks = [r_idx for r_idx, v in pairs if v is None]
                distinct = set(non_blanks)
                if len(distinct) > 1:
                    diffs_found.append((st, f, Counter(non_blanks)))
                if len(blanks) > 0 and len(non_blanks) > 0:
                    blanks_found.append((st, f, len(blanks), len(non_blanks), Counter(non_blanks)))
                elif len(non_blanks) == 0:
                    blanks_found.append((st, f, len(blanks), 0, 'ALL_BLANK'))

        print(f'Sheet [{sname}]: Total styles = {len(by_style)}')
        if diffs_found:
            print('  DIFFERENCES WITHIN SAME STYLE:')
            for item in diffs_found:
                print(f'    Style {item[0]}: field "{item[1]}" has counts: {dict(item[2])}')
        if blanks_found:
            print('  BLANKS FOUND (summary):')
            # Group by field
            field_blank_counts = Counter(item[1] for item in blanks_found)
            for f_name, count in field_blank_counts.items():
                print(f'    Field "{f_name}": {count} styles have blanks')
