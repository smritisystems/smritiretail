"""
Project      : SMRITI Retail OS
Repository   : SMRITIRetailNX
Organization : AITDL NETWORKS

Founders
- Pushpa Devi Jawahar Mallah (Founder & Chairperson)
- Jawahar Ramkripal Mallah (Founder, CEO & Chief Software Architect)

Websites     : smritibooks.com | erpnbook.com | aitdl.com
Version      : 6.92.3
Created      : 2026-10-19
Modified     : 2026-10-19
Copyright    : © AITDL.com and SMRITIBooks.com. All Rights Reserved.
License      : Proprietary Commercial Software
Classification: Internal
"""

import os
import sys
import json
import re
import openpyxl
from collections import defaultdict, Counter

sys.stdout.reconfigure(encoding="utf-8")

EXCEL_PATH = r"F:\Smriti-Clients Data\ItemmasterExported\FullItemmaster.xlsx"
CATALOG_TS = r"F:\SMRITRretailNX\src\services\unifiedFieldCatalog.ts"

def analyze_headers_and_yellow():
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    ws = wb["Item Master"]
    
    headers = [str(c.value).strip() if c.value is not None else f"COL_{idx}" for idx, c in enumerate(ws[1])]
    total_rows = ws.max_row - 1
    
    yellow_headers = []
    for idx, c in enumerate(ws[1]):
        if c.fill and c.fill.fgColor and c.fill.fgColor.rgb == "FFFFFF00":
            yellow_headers.append(headers[idx])
            
    print(f"Total rows in FullItemmaster: {total_rows}")
    print(f"Total columns: {len(headers)}")
    print(f"Headers highlighted in yellow ({len(yellow_headers)}): {yellow_headers}")
    
    # Read unifiedFieldCatalog.ts
    field_definitions = []
    if os.path.exists(CATALOG_TS):
        with open(CATALOG_TS, encoding="utf-8") as f:
            content = f.read()
        # Find all field objects
        matches = re.findall(r'\{\s*id:\s*"([^"]+)",\s*key:\s*"([^"]+)",\s*label:\s*"([^"]+)"', content)
        field_definitions = matches
        
    label_to_key = {label.strip().upper(): key for _, key, label in field_definitions}
    
    print("\nHeader Analysis against Unified Field Catalog:")
    for idx, h in enumerate(headers):
        is_yellow = h in yellow_headers
        h_upper = h.upper()
        matched_key = label_to_key.get(h_upper, "CUSTOM_OR_UNMAPPED")
        
        # Check blank counts
        blanks = 0
        for r_idx in range(2, ws.max_row + 1):
            v = ws.cell(row=r_idx, column=idx + 1).value
            if v is None or str(v).strip() == "" or str(v).strip().lower() == "none":
                blanks += 1
                
        pct = (blanks / total_rows) * 100 if total_rows else 0
        tag = "[YELLOW HIGHLIGHTED]" if is_yellow else ""
        print(f"  Col {idx:2}: {h:25} -> key: {matched_key:20} | Blanks: {blanks:4} ({pct:5.1f}%) {tag}")

if __name__ == "__main__":
    analyze_headers_and_yellow()
