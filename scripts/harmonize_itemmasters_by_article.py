"""
Project      : SMRITI Retail OS
Repository   : SMRITIRetailNX
Organization : AITDL NETWORKS

Founders
- Pushpa Devi Jawahar Mallah (Founder & Chairperson)
- Jawahar Ramkripal Mallah (Founder, CEO & Chief Software Architect)

Websites     : smritibooks.com | erpnbook.com | aitdl.com
Version      : 6.95.0
Created      : 2026-10-19
Modified     : 2026-10-19
Copyright    : © AITDL.com and SMRITIBooks.com. All Rights Reserved.
License      : Proprietary Commercial Software
Classification: Internal
"""

import os
import sys
import shutil
import json
import asyncio
from collections import defaultdict, Counter
import httpx
import psycopg2
import openpyxl

sys.stdout.reconfigure(encoding="utf-8")

DIR_PATH = r"F:\Smriti-Clients Data\Tattly Threads\ItemMaster"
BASE_URL = "https://tattlythreads.smritisys.com"

FIELDS = [
    "PRODUCT TAX",
    "HSN CODE",
    "GENDER",
    "VENDOR CODE",
    "PURCHASE CLASS",
    "DEPARTMENT",
    "MERCHANDISE CATEGORY",
    "Sub category",
    "HEELS",
    "UPPER MATERIAL",
    "OUTSOLE",
    "IMAGE LINK",
]

def clean_val(v):
    if v is None:
        return ""
    s = str(v).strip()
    return "" if s.lower() in ("none", "null") else s

def clean_num(v, default=0.0):
    if v is None:
        return default
    try:
        return float(str(v).replace(",", "").strip())
    except Exception:
        return default

def get_canonical_style_dict():
    """
    Scans all source workbooks and establishes authoritative canonical 
    values for all 12 target fields for every Article / Style Code.
    """
    print("================================================================================")
    print("PHASE 1: Building Canonical Article Metadata Catalog")
    print("================================================================================")

    # 1. Harvest reference from 1136 file
    fp_1136 = os.path.join(DIR_PATH, "TATTLY_THREADS_ALL_1136_BARCODES_WITH_STYLES.xlsx")
    ref_1136 = defaultdict(dict)
    if os.path.exists(fp_1136):
        wb1 = openpyxl.load_workbook(fp_1136, data_only=True)
        ws1 = wb1["Barcodes & Styles (1136)"]
        for r in ws1.iter_rows(min_row=2, values_only=True):
            st = clean_val(r[2]).upper()
            if not st:
                continue
            ref_1136[st] = {
                "PRODUCT TAX": clean_num(r[15], 5.0),
                "HSN CODE": clean_val(r[16]) or "64041990",
                "GENDER": clean_val(r[10]).upper() or "LADIES",
                "VENDOR CODE": clean_val(r[3]).upper(),
                "PURCHASE CLASS": "SIS",
                "DEPARTMENT": clean_val(r[11]).upper() or "LADIES FTW",
                "MERCHANDISE CATEGORY": clean_val(r[9]).upper() or "CHAPPAL",
                "HEELS": clean_val(r[17]).upper() or "FLAT",
                "UPPER MATERIAL": clean_val(r[18]).upper() or "SYNTHETIC",
                "OUTSOLE": clean_val(r[19]).upper() or "TPR",
            }

    # 2. Harvest observations from all sheets
    files_to_harvest = [
        ("TATTLY THREADS ITEM MASTER 2026 - Copy.xlsx", ["NEW", "Sheet2"]),
        ("OLD TO NEW.xlsx", ["Sheet1", "021026"]),
        ("OLD TO NEW (1).xlsx", ["Sheet1", "Sheet2", "021026"]),
    ]

    observations = defaultdict(lambda: defaultdict(list))

    for fname, sheet_names in files_to_harvest:
        fpath = os.path.join(DIR_PATH, fname)
        if not os.path.exists(fpath):
            continue
        wb = openpyxl.load_workbook(fpath, data_only=True)
        for sname in sheet_names:
            if sname not in wb.sheetnames:
                continue
            ws = wb[sname]
            h = [str(c.value).strip() if c.value else "" for c in ws[1]]
            h_map = {name: idx for idx, name in enumerate(h) if name}
            if "PRODUCT STYLE CODE" not in h_map:
                continue
            st_col = h_map["PRODUCT STYLE CODE"]

            for r in ws.iter_rows(min_row=2, values_only=True):
                if not any(r):
                    continue
                st = clean_val(r[st_col]).upper()
                if not st:
                    continue
                for f in FIELDS:
                    if f in h_map:
                        val = r[h_map[f]]
                        if val is not None:
                            val_str = str(val).strip()
                            if val_str and val_str.lower() not in ("none", "null", "ntr"):
                                observations[st][f].append(val_str)

    canonical_catalog = {}

    for st, f_dict in observations.items():
        base_st = st.rsplit("-", 1)[0] if ("-" in st and len(st.rsplit("-", 1)[1]) == 1) else st
        ref = ref_1136.get(st) or ref_1136.get(base_st) or {}

        canon = {}
        # 1. PRODUCT TAX
        canon["PRODUCT TAX"] = 5

        # 2. HSN CODE
        if st in ("CH-28", "SND-12"):
            canon["HSN CODE"] = "64032012"
        else:
            canon["HSN CODE"] = ref.get("HSN CODE") or "64041990"

        # 3. GENDER
        canon["GENDER"] = "LADIES"

        # 4. VENDOR CODE
        # Extract from style suffix if available (e.g. CH-01-A -> A)
        parts = st.split("-")
        if len(parts) >= 2 and len(parts[-1]) == 1 and parts[-1].isalpha():
            canon["VENDOR CODE"] = parts[-1].upper()
        elif f_dict.get("VENDOR CODE"):
            canon["VENDOR CODE"] = max(set(f_dict["VENDOR CODE"]), key=f_dict["VENDOR CODE"].count).upper()
        elif ref.get("VENDOR CODE"):
            canon["VENDOR CODE"] = ref["VENDOR CODE"].upper()
        elif st in ("CH-28",):
            canon["VENDOR CODE"] = "O"
        elif st in ("SND-12", "CH-35"):
            canon["VENDOR CODE"] = "P"
        elif st.startswith("CH-5"):
            canon["VENDOR CODE"] = "TT"
        else:
            canon["VENDOR CODE"] = "A"

        # 5. PURCHASE CLASS
        canon["PURCHASE CLASS"] = "SIS"

        # 6. DEPARTMENT
        canon["DEPARTMENT"] = "LADIES FTW"

        # 7. MERCHANDISE CATEGORY
        if st.startswith("CH-"):
            canon["MERCHANDISE CATEGORY"] = "CHAPPAL"
        elif st.startswith("SND-"):
            canon["MERCHANDISE CATEGORY"] = "SANDAL"
        elif st.startswith("SH-"):
            canon["MERCHANDISE CATEGORY"] = "SHOES"
        elif f_dict.get("MERCHANDISE CATEGORY"):
            cat_val = max(set(f_dict["MERCHANDISE CATEGORY"]), key=f_dict["MERCHANDISE CATEGORY"].count).upper()
            canon["MERCHANDISE CATEGORY"] = cat_val if cat_val != "OTHERS" else "CHAPPAL"
        elif ref.get("MERCHANDISE CATEGORY"):
            canon["MERCHANDISE CATEGORY"] = ref["MERCHANDISE CATEGORY"].upper()
        else:
            canon["MERCHANDISE CATEGORY"] = "CHAPPAL"

        # 8. Sub category
        # Known specific resolution: CH-13-C -> SLIP ON (Row 197 had burmy)
        if st == "CH-13-C":
            canon["Sub category"] = "SLIP ON"
        elif f_dict.get("Sub category"):
            sub_val = max(set(f_dict["Sub category"]), key=f_dict["Sub category"].count).upper()
            canon["Sub category"] = sub_val if sub_val != "OTHERS" else "STANDARD"
        else:
            canon["Sub category"] = "STANDARD"

        # 9. HEELS
        # Known specific resolution: CH-03-A -> FLAT (Row 43 had small platform)
        if st == "CH-03-A":
            canon["HEELS"] = "FLAT"
        elif f_dict.get("HEELS"):
            heel_val = max(set(f_dict["HEELS"]), key=f_dict["HEELS"].count).upper()
            canon["HEELS"] = heel_val if heel_val != "OTHERS" else (ref.get("HEELS") or "FLAT")
        elif ref.get("HEELS"):
            canon["HEELS"] = ref["HEELS"].upper()
        else:
            canon["HEELS"] = "FLAT"

        # 10. UPPER MATERIAL
        if f_dict.get("UPPER MATERIAL"):
            up_val = max(set(f_dict["UPPER MATERIAL"]), key=f_dict["UPPER MATERIAL"].count).upper()
            canon["UPPER MATERIAL"] = up_val if up_val != "OTHERS" else (ref.get("UPPER MATERIAL") or "SYNTHETIC")
        elif ref.get("UPPER MATERIAL"):
            canon["UPPER MATERIAL"] = ref["UPPER MATERIAL"].upper()
        else:
            canon["UPPER MATERIAL"] = "SYNTHETIC"

        # 11. OUTSOLE
        if f_dict.get("OUTSOLE"):
            sole_val = max(set(f_dict["OUTSOLE"]), key=f_dict["OUTSOLE"].count).upper()
            canon["OUTSOLE"] = sole_val if sole_val != "OTHERS" else (ref.get("OUTSOLE") or "TPR")
        elif ref.get("OUTSOLE"):
            canon["OUTSOLE"] = ref["OUTSOLE"].upper()
        else:
            canon["OUTSOLE"] = "TPR"

        # 12. IMAGE LINK
        canon["IMAGE LINK"] = st

        canonical_catalog[st] = canon

    print(f"  Established canonical definitions for {len(canonical_catalog)} unique styles.")
    return canonical_catalog

def harmonize_workbook(fpath, canonical_catalog):
    """
    Backs up and harmonizes all blank or differing fields for target sheets.
    """
    if not os.path.exists(fpath):
        return

    fname = os.path.basename(fpath)
    backup_path = fpath.replace(".xlsx", ".backup.xlsx")
    if not os.path.exists(backup_path):
        shutil.copy2(fpath, backup_path)
        print(f"  Created backup: {os.path.basename(backup_path)}")

    print(f"\nProcessing workbook: {fname}...")
    wb = openpyxl.load_workbook(fpath)
    total_blanks_filled = 0
    total_diffs_updated = 0

    for sname in wb.sheetnames:
        ws = wb[sname]
        h = [str(c.value).strip() if c.value else "" for c in ws[1]]
        h_map = {name: idx for idx, name in enumerate(h) if name}
        if "PRODUCT STYLE CODE" not in h_map:
            continue

        st_col = h_map["PRODUCT STYLE CODE"]
        target_cols = {f: h_map[f] for f in FIELDS if f in h_map}

        sheet_blanks = 0
        sheet_diffs = 0

        for row_idx in range(2, ws.max_row + 1):
            cell_st = ws.cell(row=row_idx, column=st_col + 1).value
            st = clean_val(cell_st).upper()
            if not st:
                continue

            canon = canonical_catalog.get(st)
            if not canon:
                base_st = st.rsplit("-", 1)[0] if ("-" in st and len(st.rsplit("-", 1)[1]) == 1) else st
                canon = canonical_catalog.get(base_st)
            if not canon:
                continue

            for f, col_idx in target_cols.items():
                cell = ws.cell(row=row_idx, column=col_idx + 1)
                curr_val = cell.value
                curr_clean = clean_val(curr_val)
                target_val = canon[f]

                # Numerical comparison for PRODUCT TAX
                if f == "PRODUCT TAX":
                    curr_num = clean_num(curr_val, None)
                    if curr_num is None:
                        cell.value = 5
                        sheet_blanks += 1
                    elif curr_num != 5.0:
                        cell.value = 5
                        sheet_diffs += 1
                elif f == "HSN CODE":
                    if not curr_clean:
                        cell.value = str(target_val)
                        sheet_blanks += 1
                    elif curr_clean != str(target_val):
                        cell.value = str(target_val)
                        sheet_diffs += 1
                else:
                    if not curr_clean or curr_clean.lower() in ("ntr", "none", "null"):
                        cell.value = target_val
                        sheet_blanks += 1
                    elif curr_clean.upper() != str(target_val).upper():
                        # Update differing field
                        cell.value = target_val
                        sheet_diffs += 1

        print(f"  [{sname:10s}] Filled blanks: {sheet_blanks:4d} | Harmonized diffs: {sheet_diffs:4d}")
        total_blanks_filled += sheet_blanks
        total_diffs_updated += sheet_diffs

    wb.save(fpath)
    print(f"  --> Saved {fname} successfully! (Total Blanks Filled: {total_blanks_filled}, Diffs Updated: {total_diffs_updated})")

def update_local_postgresql(canonical_catalog):
    print("\n================================================================================")
    print("PHASE 3: Harmonizing Local PostgreSQL Database (localhost:2781/smritisys)")
    print("================================================================================")
    try:
        conn = psycopg2.connect("postgresql://postgres:postgres@localhost:2781/smritisys")
        cur = conn.cursor()
        cur.execute("SELECT id, code, barcode, sku, style_code, gst_percentage, hsn_code, vendor_code, category, attributes FROM products;")
        rows = cur.fetchall()
        print(f"  Fetched {len(rows)} products from local PostgreSQL.")

        updated_count = 0
        for pid, code, barcode, sku, style_code, gst_pct, hsn, vc, cat, attrs in rows:
            # Determine style
            st = clean_val(style_code).upper()
            if not st:
                # Deduce style from code or sku (e.g. CH-01-A-BLACK-36 -> CH-01-A)
                target_code = (code or sku or "").upper()
                parts = target_code.split("-")
                if len(parts) >= 3 and len(parts[2]) == 1:
                    st = f"{parts[0]}-{parts[1]}-{parts[2]}"
                elif len(parts) >= 2:
                    st = f"{parts[0]}-{parts[1]}"
                elif parts:
                    st = parts[0]

            canon = canonical_catalog.get(st)
            if not canon and "-" in st:
                base_st = st.rsplit("-", 1)[0]
                canon = canonical_catalog.get(base_st)

            if not canon:
                continue

            curr_attrs = dict(attrs or {})
            changed = False

            # 1. Update columns
            new_gst = 5.0
            new_hsn = canon["HSN CODE"]
            new_vc = canon["VENDOR CODE"]
            new_cat = canon["MERCHANDISE CATEGORY"]

            update_cols = {}
            if gst_pct != 5.0:
                update_cols["gst_percentage"] = 5.0
            if clean_val(hsn) != new_hsn:
                update_cols["hsn_code"] = new_hsn
            if new_vc and clean_val(vc) != new_vc:
                update_cols["vendor_code"] = new_vc
            if new_cat and clean_val(cat).upper() != new_cat:
                update_cols["category"] = new_cat
            if not style_code or style_code != st:
                update_cols["style_code"] = st

            # 2. Update attributes
            attr_mappings = {
                "heel_type": canon["HEELS"],
                "upper_material": canon["UPPER MATERIAL"],
                "outsole": canon["OUTSOLE"],
                "outsole_material": canon["OUTSOLE"],
                "merchandise_category": canon["MERCHANDISE CATEGORY"],
                "sub_category": canon["Sub category"],
                "design_attribute": canon["Sub category"],
                "department": canon["DEPARTMENT"],
                "gender": canon["GENDER"],
                "purchase_class": canon["PURCHASE CLASS"],
                "image_link": canon["IMAGE LINK"],
                "vendor_code": canon["VENDOR CODE"],
            }

            for k, val in attr_mappings.items():
                if val and curr_attrs.get(k) != val:
                    curr_attrs[k] = val
                    changed = True

            if update_cols or changed:
                set_clauses = ["attributes = %s"]
                params = [json.dumps(curr_attrs)]
                for col_name, col_val in update_cols.items():
                    set_clauses.append(f"{col_name} = %s")
                    params.append(col_val)
                params.append(pid)

                sql = f"UPDATE products SET {', '.join(set_clauses)} WHERE id = %s;"
                cur.execute(sql, tuple(params))
                updated_count += 1

        conn.commit()
        conn.close()
        print(f"  --> Local PostgreSQL updated successfully! ({updated_count} products synchronized)")
    except Exception as e:
        print(f"  [ERROR] PostgreSQL update error: {e}")

async def update_remote_server(canonical_catalog):
    print("\n================================================================================")
    print("PHASE 4: Harmonizing Remote Production Server https://tattlythreads.smritisys.com")
    print("================================================================================")
    async with httpx.AsyncClient(timeout=120.0) as client:
        login_res = await client.post(
            f"{BASE_URL}/api/v1/auth/login",
            json={"username": "admin", "password": "Admin@123"},
        )
        if login_res.status_code != 200:
            print(f"[FATAL] Remote login failed: {login_res.status_code} {login_res.text}")
            return
        token = login_res.json()["access_token"]
        headers = {"Authorization": f"Bearer {token}", "Content-Type": "application/json"}
        print("  Successfully authenticated as Admin.")

        # 1. Fetch all existing remote products
        print("  Scanning all existing remote products...")
        all_remote = []
        page = 1
        while True:
            resp = await client.get(
                f"{BASE_URL}/api/v1/inventory/?page={page}&page_size=100",
                headers=headers,
            )
            data = resp.json()
            items = data.get("items", [])
            if not items:
                break
            all_remote.extend(items)
            if not data.get("has_next") or len(all_remote) >= data.get("total", 0):
                break
            page += 1

        print(f"  --> Found {len(all_remote)} products on remote server.")

        # Fetch and sync registered master styles
        style_res = await client.get(f"{BASE_URL}/api/v1/lookup/style_article/values", headers=headers)
        reg_styles = set()
        if style_res.status_code == 200:
            reg_styles = {
                x.get("code")
                for x in style_res.json()
                if x.get("code")
            }
        print(f"  Initially found {len(reg_styles)} registered styles in style_article.")

        # Identify all styles needed by remote products and canonical catalog
        needed_styles = set(canonical_catalog.keys())
        for p in all_remote:
            c = clean_val(p.get("code")).upper()
            sk = clean_val(p.get("sku")).upper()
            st = clean_val(p.get("style_code")).upper()
            if st:
                needed_styles.add(st)
            parts = (c or sk).split("-")
            if len(parts) >= 3 and len(parts[2]) == 1:
                needed_styles.add(f"{parts[0]}-{parts[1]}-{parts[2]}")
            elif len(parts) >= 2:
                needed_styles.add(f"{parts[0]}-{parts[1]}")
            elif parts:
                needed_styles.add(parts[0])

        missing_styles = sorted([s for s in needed_styles if s and s not in reg_styles])
        print(f"  Registering {len(missing_styles)} missing styles into style_article lookup...")
        for st in missing_styles:
            reg_res = await client.post(
                f"{BASE_URL}/api/v1/lookup/style_article/values",
                headers=headers,
                json={"code": st, "name": st, "active": True, "data": {}},
            )
            if reg_res.status_code in (200, 201):
                reg_styles.add(st)

        # Also register missing vendor codes if any
        vc_res = await client.get(f"{BASE_URL}/api/v1/lookup/vendor_code/values", headers=headers)
        reg_vcs = set()
        if vc_res.status_code == 200:
            reg_vcs = {x.get("code") for x in vc_res.json() if x.get("code")}
        needed_vcs = {c.get("VENDOR CODE") for c in canonical_catalog.values() if c.get("VENDOR CODE")}
        needed_vcs.add("X")
        needed_vcs.add("TT")
        for vc in needed_vcs:
            if vc and vc not in reg_vcs:
                await client.post(
                    f"{BASE_URL}/api/v1/lookup/vendor_code/values",
                    headers=headers,
                    json={"code": vc, "name": vc, "active": True, "data": {}},
                )
                reg_vcs.add(vc)

        print(f"  All master lookups synced ({len(reg_styles)} styles, {len(reg_vcs)} vendor codes active).")

        # 2. Update all remote products with canonical attributes
        mutated_count = 0
        skipped_count = 0
        failed_count = 0
        semaphore = asyncio.Semaphore(15)

        async def update_single(prod):
            nonlocal mutated_count, skipped_count, failed_count
            code = clean_val(prod.get("code")).upper()
            sku = clean_val(prod.get("sku")).upper()
            style_code = clean_val(prod.get("style_code")).upper()

            st = style_code
            if not st:
                parts = (code or sku).split("-")
                if len(parts) >= 3 and len(parts[2]) == 1:
                    st = f"{parts[0]}-{parts[1]}-{parts[2]}"
                elif len(parts) >= 2:
                    st = f"{parts[0]}-{parts[1]}"
                elif parts:
                    st = parts[0]

            canon = canonical_catalog.get(st)
            if not canon and "-" in st:
                base_st = st.rsplit("-", 1)[0]
                canon = canonical_catalog.get(base_st)

            if not canon:
                skipped_count += 1
                return

            curr_attrs = dict(prod.get("attributes") or {})
            changed = False

            attr_mappings = {
                "heel_type": canon["HEELS"],
                "upper_material": canon["UPPER MATERIAL"],
                "outsole": canon["OUTSOLE"],
                "outsole_material": canon["OUTSOLE"],
                "merchandise_category": canon["MERCHANDISE CATEGORY"],
                "sub_category": canon["Sub category"],
                "design_attribute": canon["Sub category"],
                "department": canon["DEPARTMENT"],
                "gender": canon["GENDER"],
                "purchase_class": canon["PURCHASE CLASS"],
                "image_link": canon["IMAGE LINK"],
                "vendor_code": canon["VENDOR CODE"],
            }

            for k, val in attr_mappings.items():
                if val and curr_attrs.get(k) != val:
                    curr_attrs[k] = val
                    changed = True

            put_body = {"attributes": curr_attrs}
            if st in reg_styles:
                put_body["style_code"] = st

            if canon.get("HSN CODE") and clean_val(prod.get("hsn_code")) != canon["HSN CODE"]:
                put_body["hsn_code"] = canon["HSN CODE"]
                changed = True
            if canon.get("VENDOR CODE") and clean_val(prod.get("vendor_code")) != canon["VENDOR CODE"]:
                put_body["vendor_code"] = canon["VENDOR CODE"]
                changed = True

            if not changed:
                skipped_count += 1
                return

            pid = prod.get("id")
            async with semaphore:
                try:
                    put_res = await client.put(
                        f"{BASE_URL}/api/v1/inventory/{pid}",
                        headers=headers,
                        json=put_body,
                    )
                    if put_res.status_code == 200:
                        mutated_count += 1
                    else:
                        failed_count += 1
                        print(f"    Warning: PUT {pid} failed ({put_res.status_code}): {put_res.text[:150]}")
                except Exception as err:
                    failed_count += 1
                    print(f"    Error PUT {pid}: {err}")

        tasks = [update_single(p) for p in all_remote]
        await asyncio.gather(*tasks)

        print(f"  --> Remote Server Update Complete: {mutated_count} updated | {skipped_count} skipped/identical | {failed_count} failed.")

        # Final Live Verification Samples
        print("\n================================================================================")
        print("PHASE 5: Live Verification Samples from Production Server")
        print("================================================================================")
        test_barcodes = ["8904551000002", "8904551003089", "8904551003577", "8904551001962", "8904551005045"]
        for bc in test_barcodes:
            check_res = await client.get(f"{BASE_URL}/api/v1/inventory/?search={bc}", headers=headers)
            if check_res.status_code == 200:
                items = check_res.json().get("items", [])
                if items:
                    p = items[0]
                    attrs = p.get("attributes", {})
                    print(f"  Barcode {bc:13s} ({p.get('code'):20s}) => HSN: {p.get('hsn_code')} | VC: {p.get('vendor_code')} | Heel: {attrs.get('heel_type')} | Upper: {attrs.get('upper_material')} | Sole: {attrs.get('outsole_material') or attrs.get('outsole')} | Cat: {attrs.get('merchandise_category')} | SubCat: {attrs.get('sub_category')}")

def main():
    canonical_catalog = get_canonical_style_dict()

    print("\n================================================================================")
    print("PHASE 2: Harmonizing Excel Workbooks on Disk")
    print("================================================================================")
    harmonize_workbook(os.path.join(DIR_PATH, "TATTLY THREADS ITEM MASTER 2026 - Copy.xlsx"), canonical_catalog)
    harmonize_workbook(os.path.join(DIR_PATH, "OLD TO NEW.xlsx"), canonical_catalog)
    harmonize_workbook(os.path.join(DIR_PATH, "OLD TO NEW (1).xlsx"), canonical_catalog)

    update_local_postgresql(canonical_catalog)
    asyncio.run(update_remote_server(canonical_catalog))

if __name__ == "__main__":
    main()
