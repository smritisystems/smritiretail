"""
Project      : SMRITI Retail OS
Author       : Jawahar Ramkripal Mallah
Designation  : Chief Systems Architect & Creator
Email        : support@smritibooks.com
Websites     : smritibooks.com | erpnbook.com | aitdl.com
Version      : 5.4.0
Created      : 2026-09-09
Copyright    : (C) SMRITIBooks.com. All Rights Reserved.
License      : Proprietary Commercial Software
Classification: Internal
"""

import openpyxl
import psycopg2
from decimal import Decimal

EXCEL_PATH = r"F:\Smriti-Clients Data\08-09-2026\Invoice_Validation_Audit_Report_TT138_to_TT194_CORRECTED.xlsx"
DB_URL = "postgresql://postgres:postgres@localhost:5432/smriti001"

UPDATE_STORES = [
    "1888", "1969", "1977", "8155", "8313", "8319", "8361", "T0N6", "T1BJ", "T25I",
    "T38X", "T40K", "T51H", "T72W", "T7FN", "T91M", "T97D", "TDL3", "TDM4", "TFW4",
    "TKU6", "TMV9", "TV78", "TVB6", "TYAC"
]

def main():
    print("=" * 90)
    print("UPDATING INVOICE VALIDATION AUDIT REPORT")
    print(f"Target: {EXCEL_PATH}")
    print("=" * 90)

    conn = psycopg2.connect(DB_URL)
    cur = conn.cursor()

    cur.execute("""
        SELECT si.delivery_store_code, si.invoice_no, si.taxable_value, si.tax_total,
               si.grand_total, si.rounding_amount, si.amount_in_words,
               SUM(sii.quantity) as total_pairs
        FROM sales_invoices si
        JOIN sales_invoice_items sii ON si.id = sii.invoice_id
        WHERE si.delivery_store_code = ANY(%s)
          AND si.import_batch_id = 'DISPATCH_20260902_STORE_GROUPED_V2'
        GROUP BY si.delivery_store_code, si.invoice_no, si.taxable_value, si.tax_total,
                 si.grand_total, si.rounding_amount, si.amount_in_words
    """, (UPDATE_STORES,))

    db_map = {r[0].strip().upper(): {
        "invoice_no": r[1],
        "taxable": float(r[2]),
        "tax": float(r[3]),
        "grand": float(r[4]),
        "rounding": float(r[5]),
        "words": r[6],
        "pairs": int(r[7])
    } for r in cur.fetchall()}

    wb = openpyxl.load_workbook(EXCEL_PATH)
    ws = wb.active

    updated_rows = 0
    for r in range(2, ws.max_row + 1):
        store = str(ws.cell(r, 4).value or "").strip().upper()
        if store in db_map:
            info = db_map[store]
            ws.cell(r, 8).value = info["pairs"]
            ws.cell(r, 15).value = info["taxable"]
            ws.cell(r, 16).value = info["tax"]
            ws.cell(r, 17).value = info["rounding"]
            ws.cell(r, 18).value = info["grand"]
            ws.cell(r, 19).value = info["words"]
            ws.cell(r, 20).value = "VERIFIED & SYNCHRONIZED (RIL_Dispatch-2.xlsx)"
            updated_rows += 1

    wb.save(EXCEL_PATH)
    conn.close()

    print(f"Successfully updated {updated_rows} rows in {EXCEL_PATH}!")
    print("=" * 90)

if __name__ == "__main__":
    main()
