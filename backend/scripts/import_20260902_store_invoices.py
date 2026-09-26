import json
import re
import uuid
from collections import defaultdict
from datetime import date, datetime, timezone
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

import psycopg2
from openpyxl import load_workbook
from pypdf import PdfReader

DB_URL = "postgresql://postgres:postgres@localhost:5432/smriti001"
SOURCE = Path(r"F:\Smriti-Clients Data\08-09-2026\RIL_Dispatch.xlsx")
PDF_DIR = SOURCE.parent
COMPANY_ID = "COMP-001"
BRANCH_ID = "MAIN"
CUSTOMER_ID = "cust-rrl-192b561d"
BATCH_ID = "DISPATCH_20260902_STORE_GROUPED_V2"
INVOICE_DATE = date(2026, 9, 2)

STATE_NAMES = {
    "06": "Haryana", "07": "Delhi", "09": "Uttar Pradesh", "10": "Bihar",
    "16": "Tripura", "18": "Assam", "19": "West Bengal", "20": "Jharkhand",
    "27": "Maharashtra", "29": "Karnataka", "32": "Kerala", "33": "Tamil Nadu",
    "36": "Telangana", "37": "Andhra Pradesh",
}


def money(value):
    return Decimal(str(value)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def normalized_color(value):
    value = str(value).strip().upper()
    value = {"CHIKOO": "CHIKKU", "GUNMETAL": "GUNMTL"}.get(value, value)
    return value.replace(" ", "")


def pdf_buyer_details(po_number):
    pdf = PDF_DIR / f"{po_number}.pdf"
    if not pdf.exists():
        raise RuntimeError(f"Missing PO PDF for {po_number}: {pdf}")
    text = " ".join((page.extract_text() or "") for page in PdfReader(str(pdf)).pages)
    marker = "SELLER DRAFT PURCHASE ORDER"
    address = text.split(marker, 1)[0].strip() if marker in text else text[:1200].strip()
    gstins = re.findall(r"\b[0-9]{2}[A-Z]{5}[0-9]{4}[A-Z][1-9A-Z]Z[0-9A-Z]\b", text)
    buyer_gstin = next((gstin for gstin in gstins if not gstin.startswith("27AAXFT")), None)
    if not buyer_gstin:
        raise RuntimeError(f"No buyer GSTIN extracted from PO {po_number}")
    return " ".join(address.split()), buyer_gstin


def load_dispatch_groups():
    workbook = load_workbook(SOURCE, data_only=True, read_only=True)
    sheet = workbook["08-09-2026"]
    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    column = {header: index for index, header in enumerate(headers)}
    groups = defaultdict(list)
    for row in sheet.iter_rows(min_row=2, values_only=True):
        values = list(row)
        if not any(value is not None for value in values):
            continue
        store = str(values[column["STORE NAME"]]).strip().upper()
        groups[store].append(values)
    return groups, column


def main():
    groups, column = load_dispatch_groups()
    if len(groups) != 57 or sum(len(rows) for rows in groups.values()) != 800:
        raise RuntimeError(f"Unexpected source shape: stores={len(groups)} rows={sum(len(rows) for rows in groups.values())}")

    connection = psycopg2.connect(DB_URL)
    connection.autocommit = False
    cursor = connection.cursor()
    try:
        cursor.execute("""
            SELECT delivery_store_code, place_of_supply_code, delivery_gstin, customer_gstin, po_reference
            FROM sales_invoices
            WHERE invoice_no LIKE 'TT2026-2027/%' AND is_deleted = false
              AND delivery_store_code IS NOT NULL
        """)
        historical = {}
        for store, state_code, delivery_gstin, customer_gstin, po_reference in cursor.fetchall():
            historical[str(store).upper()] = {
                "state_code": state_code,
                "gstin": delivery_gstin or customer_gstin,
                "po_reference": po_reference,
            }

        # These two current SIS codes have legacy aliases or no prior invoice.
        historical["8319"] = {"state_code": "29", "gstin": "29AABCR1718E1ZL", "po_reference": "5182778158"}
        historical["TV78"] = {"state_code": "06", "gstin": "06AABCR1718E1ZT", "po_reference": "5182778210"}

        cursor.execute("SELECT invoice_no FROM sales_invoices WHERE import_batch_id = %s", (BATCH_ID,))
        if cursor.fetchone():
            raise RuntimeError(f"Import batch already exists: {BATCH_ID}")

        target_numbers = [f"TT2026-2027/{number}" for number in range(138, 195)]
        cursor.execute("SELECT invoice_no FROM sales_invoices WHERE invoice_no = ANY(%s)", (target_numbers,))
        existing = [row[0] for row in cursor.fetchall()]
        if existing:
            raise RuntimeError(f"Target invoice numbers already exist: {existing}")

        cursor.execute("""
            SELECT sku, id, name, code, price, mrp, gst_percentage, hsn_code,
                   style_code, color, size
            FROM products WHERE is_deleted = false
        """)
        products = {str(row[0]).upper(): row for row in cursor.fetchall() if row[0]}

        preflight = []
        for store in sorted(groups):
            if store not in historical:
                raise RuntimeError(f"No historical tax metadata for SIS/store {store}")
            po = historical[store]["po_reference"]
            if not po:
                raise RuntimeError(f"No PO reference for SIS/store {store}")
            address, pdf_gstin = pdf_buyer_details(po)
            if not address:
                raise RuntimeError(f"No buyer address extracted from PO {po}")
            state_code = str(historical[store].get("state_code") or pdf_gstin[:2])
            gstin = historical[store].get("gstin") or pdf_gstin
            historical[store].update({"state_code": state_code, "gstin": gstin})
            preflight.append((store, po, address))

        created_products = 0
        invoice_count = 0
        line_count = 0
        quantity_total = Decimal("0")
        grand_total = Decimal("0")
        now = datetime.now(timezone.utc)

        for sequence, (store, po_reference, buyer_address) in enumerate(preflight, start=138):
            invoice_number = f"TT2026-2027/{sequence}"
            state_code = str(historical[store]["state_code"])
            gstin = historical[store]["gstin"]
            if not state_code or not gstin:
                raise RuntimeError(f"Missing GST metadata for {store}")
            lines = []
            for values in groups[store]:
                article = str(values[column["ARTICLE"]]).strip().upper()
                color = normalized_color(values[column["COLOR"]])
                mrp = money(values[column["MRP"]])
                for size in (36, 37, 38, 39, 40, 41, 42):
                    raw_quantity = values[column[size]]
                    if raw_quantity in (None, "", "-") or float(raw_quantity) <= 0:
                        continue
                    quantity = Decimal(str(raw_quantity))
                    sku = f"{article}-{color}-{size}".upper()
                    product = products.get(sku)
                    if product is None:
                        if not (article == "CH-25-G" and color == "GOLD"):
                            raise RuntimeError(f"Unresolved product master: {sku}")
                        product_id = f"prod-{uuid.uuid4().hex[:12]}"
                        product_name = f"Tattly Footwear {article} {color} Size {size}"
                        cursor.execute("""
                            INSERT INTO products (
                                code, name, price, stock, category, is_favorite, barcode,
                                brand, color, size, mrp, gst_percentage, style_code,
                                cost_price, sku, hsn_code, reserved_stock, id, uuid,
                                company_id, branch_id, created_at, modified_at,
                                is_active, is_deleted, version
                            ) VALUES (%s, %s, %s, 0, 'Footwear', false, %s,
                                      'Tattly', %s, %s, %s, 5.00, %s, 0, %s,
                                      '64032012', 0, %s, %s, %s, %s, %s, %s,
                                      true, false, 1)
                        """, (
                            sku, product_name, mrp, f"BAR-{uuid.uuid4().hex[:12]}",
                            color, str(size), mrp, article, sku, product_id,
                            str(uuid.uuid4()), COMPANY_ID, BRANCH_ID, now, now,
                        ))
                        product = (sku, product_id, product_name, sku, mrp, mrp, Decimal("5.00"), "64032012", article, color, str(size))
                        products[sku] = product
                        created_products += 1

                    taxable_value = money((mrp / Decimal("1.05")) * quantity)
                    tax_amount = money(taxable_value * Decimal("0.05"))
                    lines.append({
                        "sku": sku, "product_id": product[1], "name": product[2],
                        "quantity": quantity, "price": money(mrp / Decimal("1.05")),
                        "mrp": mrp, "taxable": taxable_value,
                        "tax": tax_amount, "total": taxable_value + tax_amount,
                    })

            taxable_total = sum((line["taxable"] for line in lines), Decimal("0"))
            tax_total = sum((line["tax"] for line in lines), Decimal("0"))
            invoice_total = taxable_total + tax_total
            invoice_id = f"inv-dispatch-{uuid.uuid4().hex[:12]}"
            snapshot = {
                "store_code": store, "sis_code": store, "po_reference": po_reference,
                "buyer_address": buyer_address, "state_code": state_code,
                "state_name": STATE_NAMES.get(state_code), "gstin": gstin,
                "source_pdf": f"{po_reference}.pdf",
            }
            cursor.execute("""
                INSERT INTO sales_invoices (
                    invoice_no, date, customer_id, tax_total, grand_total,
                    is_interstate, payment_mode, status, id, uuid, company_id,
                    branch_id, created_at, modified_at, is_active, is_deleted,
                    version, source_type, source_system, source_file, import_batch_id,
                    imported_at, import_validation_status, import_validation_notes,
                    sis_code, pos_state, po_reference, customer_name, customer_gstin,
                    billing_address, shipping_address, site_name, taxable_value,
                    rounding_amount, rule_snapshots, discount_amount, net_amount,
                    paid_amount, balance_amount, delivery_store_code, delivery_gstin,
                    delivery_location_snapshot, place_of_supply_code, billing_store_code,
                    customer_po_number_snapshot, source_document_type, source_document_id
                ) VALUES (
                    %s, %s, %s, %s, %s, %s, 'BANK_TRANSFER', 'COMPLETED', %s, %s,
                    %s, %s, %s, %s, true, false, 1, 'HISTORICAL_IMPORT',
                    'RIL_DISPATCH_XLSX', %s, %s, %s, 'VALIDATED', %s, %s, %s, %s,
                    'Reliance Retail Limited', %s, %s, %s, %s, %s, 0, %s, 0, %s,
                    0, 0, %s, %s, %s, %s, %s, %s, 'DISPATCH', 'RIL_Dispatch.xlsx'
                )
            """, (
                invoice_number, INVOICE_DATE, CUSTOMER_ID, tax_total, invoice_total,
                state_code != "27", invoice_id, str(uuid.uuid4()), COMPANY_ID, BRANCH_ID,
                now, now, str(SOURCE), BATCH_ID, now,
                f"Store-grouped dispatch import; SIS Code={store}; PO/reference={po_reference}; packing date and carton ignored.",
                store, STATE_NAMES.get(state_code), po_reference, gstin, buyer_address,
                buyer_address, store, taxable_total,
                json.dumps({"tax_rate": "5.00", "tax_inclusive_mrp": True, "grouping": "STORE_NAME", "po_reference": po_reference}),
                invoice_total, store, gstin, json.dumps(snapshot), state_code, store, po_reference,
            ))

            for line_number, line in enumerate(lines, start=1):
                interstate = state_code != "27"
                igst = line["tax"] if interstate else Decimal("0")
                cgst = Decimal("0") if interstate else money(line["tax"] / 2)
                sgst = Decimal("0") if interstate else money(line["tax"] / 2)
                cursor.execute("""
                    INSERT INTO sales_invoice_items (
                        invoice_id, product_id, code, name, quantity, price, hsn_code,
                        gst_rate, tax_amount, total_amount, mrp, disc_pct, taxable_value,
                        igst_amount, cgst_amount, sgst_amount, line_no,
                        source_line_type, source_line_id
                    ) VALUES (%s, %s, %s, %s, %s, %s, '64032012', 5.00, %s, %s,
                              %s, 0, %s, %s, %s, %s, %s, 'DISPATCH_IMPORT', %s)
                """, (
                    invoice_id, line["product_id"], line["sku"], line["name"],
                    line["quantity"], line["price"], line["tax"], line["total"],
                    line["mrp"], line["taxable"], igst, cgst, sgst, line_number,
                    f"{store}:{line_number}",
                ))
            invoice_count += 1
            line_count += len(lines)
            quantity_total += sum((line["quantity"] for line in lines), Decimal("0"))
            grand_total += invoice_total

        connection.commit()
        print(json.dumps({
            "status": "SUCCESS", "batch": BATCH_ID, "invoice_count": invoice_count,
            "first_invoice": "TT2026-2027/138", "last_invoice": "TT2026-2027/194",
            "line_count": line_count, "quantity": str(quantity_total),
            "grand_total": str(grand_total), "created_products": created_products,
        }, indent=2))
    except Exception:
        connection.rollback()
        raise
    finally:
        cursor.close()
        connection.close()


if __name__ == "__main__":
    main()
