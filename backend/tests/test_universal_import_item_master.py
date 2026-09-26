"""
Project      : SMRITI Retail OS
Author       : Jawahar Ramkripal Mallah
Designation  : Chief Systems Architect & Creator
Email        : support@smritibooks.com
Websites     : smritibooks.com | erpnbook.com | aitdl.com
Version      : 2.1.0
Created      : 2026-09-26
Copyright    : © SMRITIBooks.com. All Rights Reserved.
License      : Proprietary Commercial Software
Classification: Internal Test Suite
"""

import sys
import os
import uuid
from io import BytesIO
from pathlib import Path
from decimal import Decimal

backend_dir = Path(__file__).resolve().parent.parent
if str(backend_dir) not in sys.path:
    sys.path.insert(0, str(backend_dir))

from dotenv import dotenv_values
env_file = backend_dir.parent / ".env"
if env_file.exists():
    for k, v in dotenv_values(env_file).items():
        if v is not None and k not in os.environ:
            os.environ[k] = v

import pytest
import openpyxl
from fastapi import HTTPException
from sqlalchemy import select, delete
from app.db.session import get_company_sessionmaker
from app.models.item_master import Item, ItemVariant, ItemBarcode, ItemWarehouseLocation
from app.models.pricing import PriceBook, PriceBookEntry
from app.api.v1.universal_import import (
    preview_universal_import,
    commit_universal_import,
    download_item_master_template,
    ImportPreviewRequest,
    ImportCommitRequest,
)

TEST_STYLES = [
    "CH-TEST-STYLE-01",
    "CH-CASCADE-01",
    "CH-CONFLICT-01",
    "CH-PRICEMODE-01",
    "CH-MATRIX-01",
    "CH-UPDATE-01",
]
TEST_BARCODES = [f"8909999000{i:03d}" for i in range(1, 25)]


@pytest.fixture(autouse=True)
async def cleanup_test_data():
    session_factory = get_company_sessionmaker("smriti001")
    async with session_factory() as session:
        await session.execute(delete(ItemBarcode).where(ItemBarcode.barcode.in_(TEST_BARCODES)))
        sub = select(Item.id).where(Item.item_code.in_(TEST_STYLES))
        await session.execute(delete(PriceBookEntry).where(PriceBookEntry.item_id.in_(sub)))
        await session.execute(delete(ItemWarehouseLocation).where(ItemWarehouseLocation.item_id.in_(sub)))
        await session.execute(delete(ItemBarcode).where(ItemBarcode.item_id.in_(sub)))
        await session.execute(delete(ItemVariant).where(ItemVariant.item_id.in_(sub)))
        await session.execute(delete(Item).where(Item.item_code.in_(TEST_STYLES)))
        await session.commit()
    yield
    async with session_factory() as session:
        await session.execute(delete(ItemBarcode).where(ItemBarcode.barcode.in_(TEST_BARCODES)))
        sub = select(Item.id).where(Item.item_code.in_(TEST_STYLES))
        await session.execute(delete(PriceBookEntry).where(PriceBookEntry.item_id.in_(sub)))
        await session.execute(delete(ItemWarehouseLocation).where(ItemWarehouseLocation.item_id.in_(sub)))
        await session.execute(delete(ItemBarcode).where(ItemBarcode.item_id.in_(sub)))
        await session.execute(delete(ItemVariant).where(ItemVariant.item_id.in_(sub)))
        await session.execute(delete(Item).where(Item.item_code.in_(TEST_STYLES)))
        await session.commit()


@pytest.mark.asyncio
async def test_item_master_dry_run_preview_valid():
    """Verify that preview dry-run correctly reconciles valid multi-variant style rows as NEW."""
    session_factory = get_company_sessionmaker("smriti001")
    async with session_factory() as session:
        req = ImportPreviewRequest(
            target="ITEM_MASTER",
            rows=[
                {
                    "rowNumber": 1,
                    "BARCODE_NO": TEST_BARCODES[0],
                    "ARTICLE_STYLE_CODE": TEST_STYLES[0],
                    "BRAND_NAME": "TATTLY THREADS",
                    "COLOR": "BLACK",
                    "SIZE": "36",
                    "MRP": 1899,
                    "SELLING_PRICE": 1899,
                    "BUYING_PRICE": 470,
                    "LANDED_COST_PRICE": 495,
                    "WAREHOUSE_CODE": "WH-MAIN",
                    "REORDER_LEVEL": 10,
                },
                {
                    "rowNumber": 2,
                    "BARCODE_NO": TEST_BARCODES[1],
                    "ARTICLE_STYLE_CODE": TEST_STYLES[0],
                    "BRAND_NAME": "TATTLY THREADS",
                    "COLOR": "BLACK",
                    "SIZE": "37",
                    "MRP": 1899,
                    "SELLING_PRICE": 1899,
                    "BUYING_PRICE": 470,
                    "LANDED_COST_PRICE": 495,
                    "WAREHOUSE_CODE": "WH-MAIN",
                    "REORDER_LEVEL": 10,
                },
            ]
        )
        res = await preview_universal_import(request=req, db=session, _current_user={"company_id": "COMP-001"})
        assert res["target"] == "ITEM_MASTER"
        assert res["counts"]["total"] == 2
        assert res["counts"]["valid"] == 2
        assert res["counts"]["new"] == 2
        assert res["counts"]["distinct_styles"] == 1
        assert res["summary"]["status"] == "READY_FOR_IMPORT"
        assert len(res["rows"]) == 2
        assert res["rows"][0]["status"] == "VALID"
        assert res["rows"][0]["reconciliation_state"] == "NEW"
        assert res["rows"][1]["action"] == "ATTACH_VARIANT_TO_STYLE"
        assert res["rows"][1]["reconciliation_state"] == "NEW"


@pytest.mark.asyncio
async def test_item_master_dry_run_detects_conflicts():
    """Verify duplicate in file and pricing violation detection in pre-flight dry-run."""
    session_factory = get_company_sessionmaker("smriti001")
    async with session_factory() as session:
        req = ImportPreviewRequest(
            target="ITEM_MASTER",
            rows=[
                {
                    "rowNumber": 1,
                    "BARCODE_NO": TEST_BARCODES[0],
                    "ARTICLE_STYLE_CODE": TEST_STYLES[0],
                    "COLOR": "BLACK",
                    "SIZE": "36",
                    "MRP": 1899,
                    "SELLING_PRICE": 1899,
                },
                {
                    "rowNumber": 2,
                    "BARCODE_NO": TEST_BARCODES[0],  # Duplicate barcode in batch!
                    "ARTICLE_STYLE_CODE": TEST_STYLES[0],
                    "COLOR": "BLACK",
                    "SIZE": "37",
                    "MRP": 1899,
                    "SELLING_PRICE": 2100,  # Selling price > MRP!
                },
            ]
        )
        res = await preview_universal_import(request=req, db=session, _current_user={"company_id": "COMP-001"})
        assert res["counts"]["duplicate_in_file"] >= 1
        assert res["counts"]["pricing_conflicts"] >= 1
        assert res["summary"]["status"] == "VALIDATION_ISSUES_FOUND"
        assert res["rows"][1]["status"] == "INVALID"
        assert res["rows"][1]["reconciliation_state"] == "DUPLICATE_IN_FILE"


@pytest.mark.asyncio
async def test_item_master_commit_3tier_cascade():
    """Verify full 3-tier cascade commit: Item -> ItemVariant -> ItemBarcode -> WarehouseLocation -> PriceBookEntry."""
    session_factory = get_company_sessionmaker("smriti001")
    async with session_factory() as session:
        user_ctx = {"id": "usr-test", "company_id": "COMP-001", "branch_id": "BR-001"}
        idempotency_key = f"idem-{uuid.uuid4().hex[:12]}"

        commit_req = ImportCommitRequest(
            target="ITEM_MASTER",
            idempotency_key=idempotency_key,
            price_mode="CREATE_AS_DRAFT",
            rows=[
                {
                    "rowNumber": 1,
                    "BARCODE_NO": TEST_BARCODES[0],
                    "SKU_PREVIEW": f"{TEST_STYLES[1]}-BLACK-36-O",
                    "ARTICLE_STYLE_CODE": TEST_STYLES[1],
                    "BRAND_NAME": "TATTLY THREADS",
                    "MERCHANDISE_DEPARTMENT": "FOOTWEAR",
                    "MERCHANDISE_CATEGORY": "LADIES FOOTWEAR",
                    "PRODUCT_TYPE": "CHAPPAL",
                    "COLOR": "BLACK",
                    "SIZE": "36",
                    "MRP": 1899,
                    "SELLING_PRICE": 1899,
                    "BUYING_PRICE": 470,
                    "LANDED_COST_PRICE": 495,
                    "GST_RATE_PERCENT": 5,
                    "HSN_CODE": "6402",
                    "GENDER": "LADIES",
                    "WAREHOUSE_CODE": "WH-MAIN",
                    "REORDER_LEVEL": 15,
                },
                {
                    "rowNumber": 2,
                    "BARCODE_NO": TEST_BARCODES[1],
                    "SKU_PREVIEW": f"{TEST_STYLES[1]}-BLACK-37-O",
                    "ARTICLE_STYLE_CODE": TEST_STYLES[1],
                    "BRAND_NAME": "TATTLY THREADS",
                    "MERCHANDISE_DEPARTMENT": "FOOTWEAR",
                    "MERCHANDISE_CATEGORY": "LADIES FOOTWEAR",
                    "PRODUCT_TYPE": "CHAPPAL",
                    "COLOR": "BLACK",
                    "SIZE": "37",
                    "MRP": 1899,
                    "SELLING_PRICE": 1899,
                    "BUYING_PRICE": 470,
                    "LANDED_COST_PRICE": 495,
                    "GST_RATE_PERCENT": 5,
                    "HSN_CODE": "6402",
                    "GENDER": "LADIES",
                    "WAREHOUSE_CODE": "WH-MAIN",
                    "REORDER_LEVEL": 15,
                },
            ]
        )

        class MockTenant:
            id = "smriti001"
            code = "smriti001"

        res = await commit_universal_import(
            request=commit_req,
            db=session,
            current_user=user_ctx,
            tenant=MockTenant(),
        )

        assert res["success"] is True
        assert res["idempotent_replay"] is False
        assert len(res["results"]) == 2

        # 1. Exactly 1 Parent Item
        items = (await session.execute(select(Item).where(Item.item_code == TEST_STYLES[1], Item.company_id == "COMP-001"))).scalars().all()
        assert len(items) == 1
        parent = items[0]
        assert parent.brand.upper() == "TATTLY THREADS"
        assert parent.style_code == TEST_STYLES[1]

        # 2. Exactly 2 Child Variants
        variants = (await session.execute(select(ItemVariant).where(ItemVariant.item_id == parent.id))).scalars().all()
        assert len(variants) == 2
        var_skus = {v.variant_sku for v in variants}
        assert f"{TEST_STYLES[1]}-BLACK-36-O" in var_skus
        assert f"{TEST_STYLES[1]}-BLACK-37-O" in var_skus

        # 3. Barcodes: 2 variant-level barcodes matching the imported lines
        bcs = (await session.execute(select(ItemBarcode).where(ItemBarcode.item_id == parent.id))).scalars().all()
        variant_bcs = [b for b in bcs if b.variant_id is not None]
        assert len(variant_bcs) == 2
        assert {b.barcode for b in variant_bcs} == {TEST_BARCODES[0], TEST_BARCODES[1]}

        # 4. Warehouse location
        locs = (await session.execute(select(ItemWarehouseLocation).where(ItemWarehouseLocation.item_id == parent.id))).scalars().all()
        assert len(locs) == 1
        assert locs[0].warehouse_id == "WH-MAIN"
        assert float(locs[0].min_reorder_level) == 15.0

        # 5. Price Book Entry (Draft Mode)
        pbes = (await session.execute(select(PriceBookEntry).where(PriceBookEntry.item_id == parent.id))).scalars().all()
        assert len(pbes) == 2
        assert float(pbes[0].selling_price) == 1899.0
        assert float(pbes[0].mrp) == 1899.0

        # 6. Test Idempotency
        replay_res = await commit_universal_import(
            request=commit_req,
            db=session,
            current_user=user_ctx,
            tenant=MockTenant(),
        )
        assert replay_res["success"] is True
        assert replay_res["idempotent_replay"] is True


@pytest.mark.asyncio
async def test_reconciliation_states_existing_match_vs_conflict():
    """Verify distinction between EXISTING_MATCH (same product/variant) and EXISTING_CONFLICT (different product)."""
    session_factory = get_company_sessionmaker("smriti001")
    async with session_factory() as session:
        user_ctx = {"id": "usr-test", "company_id": "COMP-001", "branch_id": "BR-001"}

        class MockTenant:
            id = "smriti001"
            code = "smriti001"

        # 1. Commit initial item
        init_commit = ImportCommitRequest(
            target="ITEM_MASTER",
            idempotency_key=f"idem-{uuid.uuid4().hex[:12]}",
            price_mode="DO_NOT_CREATE",
            rows=[
                {
                    "rowNumber": 1,
                    "BARCODE_NO": TEST_BARCODES[3],
                    "ARTICLE_STYLE_CODE": TEST_STYLES[2],
                    "COLOR": "TAN",
                    "SIZE": "38",
                    "MRP": 1499,
                    "SELLING_PRICE": 1499,
                }
            ]
        )
        await commit_universal_import(request=init_commit, db=session, current_user=user_ctx, tenant=MockTenant())

        # 2. Preview same barcode with SAME style & SKU -> EXISTING_MATCH (VALID, SKIP)
        match_req = ImportPreviewRequest(
            target="ITEM_MASTER",
            rows=[
                {
                    "rowNumber": 1,
                    "BARCODE_NO": TEST_BARCODES[3],
                    "ARTICLE_STYLE_CODE": TEST_STYLES[2],
                    "COLOR": "TAN",
                    "SIZE": "38",
                    "MRP": 1499,
                    "SELLING_PRICE": 1499,
                }
            ]
        )
        match_res = await preview_universal_import(request=match_req, db=session, _current_user=user_ctx)
        assert match_res["counts"]["existing_match"] == 1
        assert match_res["counts"]["existing_conflict"] == 0
        assert match_res["rows"][0]["reconciliation_state"] == "EXISTING_MATCH"
        assert match_res["rows"][0]["status"] == "VALID"
        assert match_res["rows"][0]["action"] == "SKIP"

        # 3. Preview same barcode with DIFFERENT style -> EXISTING_CONFLICT (INVALID, BLOCK)
        conflict_req = ImportPreviewRequest(
            target="ITEM_MASTER",
            rows=[
                {
                    "rowNumber": 1,
                    "BARCODE_NO": TEST_BARCODES[3],  # Same barcode!
                    "ARTICLE_STYLE_CODE": "COMPLETELY-DIFFERENT-STYLE",  # Conflict!
                    "COLOR": "RED",
                    "SIZE": "42",
                    "MRP": 2999,
                    "SELLING_PRICE": 2999,
                }
            ]
        )
        conflict_res = await preview_universal_import(request=conflict_req, db=session, _current_user=user_ctx)
        assert conflict_res["counts"]["existing_conflict"] == 1
        assert conflict_res["rows"][0]["reconciliation_state"] == "EXISTING_CONFLICT"
        assert conflict_res["rows"][0]["status"] == "INVALID"
        assert conflict_res["rows"][0]["action"] == "BLOCK"


@pytest.mark.asyncio
async def test_commit_existing_match_modes():
    """Verify existing_match_mode options: SKIP, UPDATE_METADATA_AND_PRICE, FAIL_ON_EXISTING."""
    session_factory = get_company_sessionmaker("smriti001")
    async with session_factory() as session:
        user_ctx = {"id": "usr-test", "company_id": "COMP-001", "branch_id": "BR-001"}

        class MockTenant:
            id = "smriti001"
            code = "smriti001"

        # Initial commit
        await commit_universal_import(
            request=ImportCommitRequest(
                target="ITEM_MASTER",
                idempotency_key=f"idem-{uuid.uuid4().hex[:12]}",
                price_mode="DO_NOT_CREATE",
                rows=[
                    {
                        "rowNumber": 1,
                        "BARCODE_NO": TEST_BARCODES[5],
                        "ARTICLE_STYLE_CODE": TEST_STYLES[5],
                        "COLOR": "BLUE",
                        "SIZE": "40",
                        "MRP": 1000,
                        "SELLING_PRICE": 1000,
                    }
                ]
            ),
            db=session,
            current_user=user_ctx,
            tenant=MockTenant(),
        )

        # Mode: SKIP
        skip_res = await commit_universal_import(
            request=ImportCommitRequest(
                target="ITEM_MASTER",
                idempotency_key=f"idem-{uuid.uuid4().hex[:12]}",
                existing_match_mode="SKIP",
                rows=[
                    {
                        "rowNumber": 1,
                        "BARCODE_NO": TEST_BARCODES[5],
                        "ARTICLE_STYLE_CODE": TEST_STYLES[5],
                        "COLOR": "BLUE",
                        "SIZE": "40",
                        "MRP": 1200,
                        "SELLING_PRICE": 1200,
                    }
                ]
            ),
            db=session,
            current_user=user_ctx,
            tenant=MockTenant(),
        )
        assert skip_res["results"][0]["status"] == "SKIPPED_EXISTING_MATCH"

        # Mode: UPDATE_METADATA_AND_PRICE
        update_res = await commit_universal_import(
            request=ImportCommitRequest(
                target="ITEM_MASTER",
                idempotency_key=f"idem-{uuid.uuid4().hex[:12]}",
                existing_match_mode="UPDATE_METADATA_AND_PRICE",
                rows=[
                    {
                        "rowNumber": 1,
                        "BARCODE_NO": TEST_BARCODES[5],
                        "ARTICLE_STYLE_CODE": TEST_STYLES[5],
                        "COLOR": "BLUE",
                        "SIZE": "40",
                        "MRP": 1650,
                        "SELLING_PRICE": 1500,
                        "LANDED_COST_PRICE": 800,
                    }
                ]
            ),
            db=session,
            current_user=user_ctx,
            tenant=MockTenant(),
        )
        assert update_res["results"][0]["status"] == "UPDATED_EXISTING_MATCH"

        # Verify updated pricing in DB
        variant = (await session.execute(
            select(ItemVariant).where(ItemVariant.variant_sku == f"{TEST_STYLES[5]}-BLUE-40", ItemVariant.company_id == "COMP-001")
        )).scalars().first()
        assert float(variant.mrp) == 1650.0
        assert float(variant.selling_price) == 1500.0

        # Mode: FAIL_ON_EXISTING
        with pytest.raises(HTTPException) as exc:
            await commit_universal_import(
                request=ImportCommitRequest(
                    target="ITEM_MASTER",
                    idempotency_key=f"idem-{uuid.uuid4().hex[:12]}",
                    existing_match_mode="FAIL_ON_EXISTING",
                    rows=[
                        {
                            "rowNumber": 1,
                            "BARCODE_NO": TEST_BARCODES[5],
                            "ARTICLE_STYLE_CODE": TEST_STYLES[5],
                            "COLOR": "BLUE",
                            "SIZE": "40",
                            "MRP": 1650,
                            "SELLING_PRICE": 1500,
                        }
                    ]
                ),
                db=session,
                current_user=user_ctx,
                tenant=MockTenant(),
            )
        assert exc.value.status_code == 409


@pytest.mark.asyncio
async def test_pricing_modes_coverage():
    """Verify DO_NOT_CREATE vs CREATE_LIVE_RETAIL pricing modes."""
    session_factory = get_company_sessionmaker("smriti001")
    async with session_factory() as session:
        user_ctx = {"id": "usr-test", "company_id": "COMP-001", "branch_id": "BR-001"}

        class MockTenant:
            id = "smriti001"
            code = "smriti001"

        # 1. DO_NOT_CREATE mode
        res1 = await commit_universal_import(
            request=ImportCommitRequest(
                target="ITEM_MASTER",
                idempotency_key=f"idem-{uuid.uuid4().hex[:12]}",
                price_mode="DO_NOT_CREATE",
                rows=[
                    {
                        "rowNumber": 1,
                        "BARCODE_NO": TEST_BARCODES[6],
                        "ARTICLE_STYLE_CODE": TEST_STYLES[3],
                        "COLOR": "GREY",
                        "SIZE": "39",
                        "MRP": 999,
                        "SELLING_PRICE": 999,
                    }
                ]
            ),
            db=session,
            current_user=user_ctx,
            tenant=MockTenant(),
        )
        assert res1["success"] is True
        item1 = (await session.execute(select(Item).where(Item.item_code == TEST_STYLES[3]))).scalars().first()
        pbes1 = (await session.execute(select(PriceBookEntry).where(PriceBookEntry.item_id == item1.id))).scalars().all()
        assert len(pbes1) == 0

        # 2. CREATE_LIVE_RETAIL mode
        res2 = await commit_universal_import(
            request=ImportCommitRequest(
                target="ITEM_MASTER",
                idempotency_key=f"idem-{uuid.uuid4().hex[:12]}",
                price_mode="CREATE_LIVE_RETAIL",
                rows=[
                    {
                        "rowNumber": 1,
                        "BARCODE_NO": TEST_BARCODES[7],
                        "ARTICLE_STYLE_CODE": TEST_STYLES[3],
                        "COLOR": "GREY",
                        "SIZE": "40",
                        "MRP": 999,
                        "SELLING_PRICE": 999,
                    }
                ]
            ),
            db=session,
            current_user=user_ctx,
            tenant=MockTenant(),
        )
        assert res2["success"] is True
        pbes2 = (await session.execute(select(PriceBookEntry).where(PriceBookEntry.item_id == item1.id))).scalars().all()
        assert len(pbes2) == 1
        pb = (await session.execute(select(PriceBook).where(PriceBook.id == pbes2[0].price_book_id))).scalars().first()
        assert pb.status == "ACTIVE"


@pytest.mark.asyncio
async def test_multi_variant_matrix_colors_and_sizes():
    """Verify ingesting a 4-variant matrix (2 colors x 2 sizes) under 1 parent style."""
    session_factory = get_company_sessionmaker("smriti001")
    async with session_factory() as session:
        user_ctx = {"id": "usr-test", "company_id": "COMP-001", "branch_id": "BR-001"}

        class MockTenant:
            id = "smriti001"
            code = "smriti001"

        matrix_rows = [
            {"rowNumber": 1, "BARCODE_NO": TEST_BARCODES[8], "ARTICLE_STYLE_CODE": TEST_STYLES[4], "COLOR": "BLACK", "SIZE": "36", "MRP": 1299, "SELLING_PRICE": 1299},
            {"rowNumber": 2, "BARCODE_NO": TEST_BARCODES[9], "ARTICLE_STYLE_CODE": TEST_STYLES[4], "COLOR": "BLACK", "SIZE": "37", "MRP": 1299, "SELLING_PRICE": 1299},
            {"rowNumber": 3, "BARCODE_NO": TEST_BARCODES[10], "ARTICLE_STYLE_CODE": TEST_STYLES[4], "COLOR": "WHITE", "SIZE": "36", "MRP": 1299, "SELLING_PRICE": 1299},
            {"rowNumber": 4, "BARCODE_NO": TEST_BARCODES[11], "ARTICLE_STYLE_CODE": TEST_STYLES[4], "COLOR": "WHITE", "SIZE": "37", "MRP": 1299, "SELLING_PRICE": 1299},
        ]

        # Dry run preview check
        preview = await preview_universal_import(
            request=ImportPreviewRequest(target="ITEM_MASTER", rows=matrix_rows),
            db=session,
            _current_user=user_ctx,
        )
        assert preview["counts"]["distinct_styles"] == 1
        assert preview["counts"]["new"] == 4
        assert preview["rows"][0]["action"] == "CREATE_ITEM_AND_VARIANT"
        assert preview["rows"][1]["action"] == "ATTACH_VARIANT_TO_STYLE"
        assert preview["rows"][2]["action"] == "ATTACH_VARIANT_TO_STYLE"
        assert preview["rows"][3]["action"] == "ATTACH_VARIANT_TO_STYLE"

        # Commit check
        commit = await commit_universal_import(
            request=ImportCommitRequest(
                target="ITEM_MASTER",
                idempotency_key=f"idem-{uuid.uuid4().hex[:12]}",
                rows=matrix_rows,
            ),
            db=session,
            current_user=user_ctx,
            tenant=MockTenant(),
        )
        assert commit["success"] is True
        assert len(commit["results"]) == 4

        # Verify DB hierarchy
        parent_items = (await session.execute(select(Item).where(Item.item_code == TEST_STYLES[4]))).scalars().all()
        assert len(parent_items) == 1
        variants = (await session.execute(select(ItemVariant).where(ItemVariant.item_id == parent_items[0].id))).scalars().all()
        assert len(variants) == 4
        skus = {v.variant_sku for v in variants}
        assert skus == {
            f"{TEST_STYLES[4]}-BLACK-36",
            f"{TEST_STYLES[4]}-BLACK-37",
            f"{TEST_STYLES[4]}-WHITE-36",
            f"{TEST_STYLES[4]}-WHITE-37",
        }


@pytest.mark.asyncio
async def test_dynamic_template_generation_endpoint():
    """Verify GET /api/v1/universal-import/templates/item-master.xlsx generates a valid populated workbook."""
    session_factory = get_company_sessionmaker("smriti001")
    async with session_factory() as session:
        response = await download_item_master_template(
            company_id="COMP-001",
            warehouse_id="WH-SHOP",
            db=session,
            _current_user={"company_id": "COMP-001"},
        )
        assert response.status_code == 200
        assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in response.media_type

        # Read streaming body into openpyxl
        body = b""
        async for chunk in response.body_iterator:
            body += chunk
        
        wb = openpyxl.load_workbook(BytesIO(body))
        assert "Item Master Template" in wb.sheetnames
        assert "Validation Lists" in wb.sheetnames

        it_ws = wb["Item Master Template"]
        # Check pre-filled warehouse_id in row 5
        assert it_ws.cell(row=5, column=28).value == "WH-SHOP"

        # Check live warehouses populated in Validation Lists
        vl_ws = wb["Validation Lists"]
        wh_values = [vl_ws.cell(row=r, column=19).value for r in range(2, 10)]
        assert "WH-MAIN" in wh_values
