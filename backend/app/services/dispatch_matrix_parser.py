"""
Project      : SMRITI Retail OS
Author       : Jawahar Ramkripal Mallah
Designation  : Chief Systems Architect & Creator
Email        : support@smritibooks.com
Websites     : smritibooks.com | erpnbook.com | aitdl.com
Version      : 6.40.1
Created      : 2026-09-18
Copyright    : © SMRITIBooks.com. All Rights Reserved.
License      : Proprietary Commercial Software
Classification: Canonical Dispatch Matrix Parser Engine
"""

import io
import re
from decimal import Decimal
from typing import Dict, List, Any, Optional, Tuple
import openpyxl

# smriti_capability(entity="SALES", capability="B2B_DISPATCH_INVOICING_STUDIO", role="CORE", canonicalOwner="backend/app/services/dispatch_matrix_parser.py")


class DispatchMatrixParseError(Exception):
    """Raised when an uploaded dispatch sheet cannot be parsed."""
    pass


class DispatchMatrixParser:
    """
    Enterprise Dynamic Matrix Parser for B2B Client Dispatch Workbooks.
    
    Dynamically identifies:
    - Store identifier column ('STORE NAME', 'SITE', 'STORE', etc.)
    - Product identity columns ('ARTICLE', 'COLOR', 'MRP')
    - Arbitrary size matrix columns (e.g. 36..42, S..XXL, 6..11) between MRP and TOTAL
    - Aggregate column ('TOTAL', 'QTY', 'SUM')
    
    Unpivots horizontal size matrices into normalized line item records.
    """

    STORE_ALIASES = ["STORE NAME", "STORE", "SITE", "STORE CODE", "SITE CODE", "STORE_NAME", "SITE_NAME"]
    ARTICLE_ALIASES = ["ARTICLE", "SKU", "STYLE", "ITEM CODE", "ITEM", "PRODUCT", "STYLE CODE"]
    COLOR_ALIASES = ["COLOR", "COLOUR", "SHADE"]
    MRP_ALIASES = ["MRP", "RETAIL PRICE", "MAX RETAIL PRICE", "TAG PRICE", "PRICE"]
    TOTAL_ALIASES = ["TOTAL", "TOTAL QTY", "QTY", "SUM", "TOTAL PAIRS", "PAIRS"]

    @classmethod
    def parse_workbook(
        cls,
        file_bytes: bytes,
        sheet_name: Optional[str] = None
    ) -> Dict[str, Any]:
        """
        Parses an Excel workbook from memory, returning workbook metadata and normalized lines.
        """
        try:
            wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
        except Exception as exc:
            raise DispatchMatrixParseError(f"Invalid Excel workbook format: {str(exc)}") from exc

        available_sheets = wb.sheetnames
        if not available_sheets:
            raise DispatchMatrixParseError("Excel workbook contains no sheets.")

        target_sheet_name = sheet_name if sheet_name and sheet_name in available_sheets else available_sheets[0]
        ws = wb[target_sheet_name]

        parsed_data = cls._parse_worksheet(ws)
        return {
            "sheet_name": target_sheet_name,
            "available_sheets": available_sheets,
            "total_rows": parsed_data["total_rows"],
            "detected_sizes": parsed_data["detected_sizes"],
            "store_groups": parsed_data["store_groups"],
            "raw_lines": parsed_data["raw_lines"],
            "column_map": parsed_data["column_map"]
        }

    @classmethod
    def _find_header_row(cls, ws: Any) -> Tuple[int, Dict[str, int], List[str]]:
        """
        Scans top 15 rows to find the primary header row with matching aliases.
        """
        for r in range(1, min(ws.max_row + 1, 16)):
            row_vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
            str_vals = [str(v).strip().upper() for v in row_vals if v is not None]

            # Check if at least STORE and ARTICLE match
            has_store = any(alias in str_vals for alias in cls.STORE_ALIASES)
            has_article = any(alias in str_vals for alias in cls.ARTICLE_ALIASES)

            if has_store and has_article:
                col_map: Dict[str, int] = {}
                col_names: List[str] = []
                for c in range(1, ws.max_column + 1):
                    val = ws.cell(r, c).value
                    if val is not None:
                        val_str = str(val).strip()
                        col_map[val_str] = c
                        col_names.append(val_str)
                return r, col_map, col_names

        raise DispatchMatrixParseError(
            "Could not detect dispatch header row. Expected columns: 'STORE NAME', 'ARTICLE', 'COLOR', 'MRP'."
        )

    @classmethod
    def _resolve_canonical_column(
        cls,
        col_map: Dict[str, int],
        aliases: List[str]
    ) -> Tuple[str, int]:
        """Matches column map entries against alias list."""
        for name, col_idx in col_map.items():
            clean_name = str(name).strip().upper()
            if clean_name in aliases:
                return name, col_idx
        raise DispatchMatrixParseError(f"Missing required dispatch column. Expected one of: {aliases}")

    @classmethod
    def _parse_worksheet(cls, ws: Any) -> Dict[str, Any]:
        header_row_idx, col_map, col_names = cls._find_header_row(ws)

        store_col_name, store_col_idx = cls._resolve_canonical_column(col_map, cls.STORE_ALIASES)
        art_col_name, art_col_idx = cls._resolve_canonical_column(col_map, cls.ARTICLE_ALIASES)
        color_col_name, color_col_idx = cls._resolve_canonical_column(col_map, cls.COLOR_ALIASES)
        mrp_col_name, mrp_col_idx = cls._resolve_canonical_column(col_map, cls.MRP_ALIASES)

        # Detect TOTAL column if present
        total_col_idx: Optional[int] = None
        for name, idx in col_map.items():
            if str(name).strip().upper() in cls.TOTAL_ALIASES:
                total_col_idx = idx
                break

        # Detect Size Matrix columns
        # Size columns typically sit between MRP and TOTAL, or have numeric/apparel size values
        size_cols: List[Tuple[str, int]] = []
        for name, col_idx in col_map.items():
            if col_idx in (store_col_idx, art_col_idx, color_col_idx, mrp_col_idx, total_col_idx):
                continue
            
            clean_name = str(name).strip()
            # Numeric sizes (e.g. "36", 36, "37", "6", "7")
            if clean_name.isdigit():
                size_cols.append((clean_name, col_idx))
            # Standard apparel sizes (S, M, L, XL, XXL, etc.)
            elif clean_name.upper() in ["XS", "S", "M", "L", "XL", "XXL", "XXXL", "2XL", "3XL", "FREE", "FS"]:
                size_cols.append((clean_name.upper(), col_idx))
            # Column positioned between MRP and TOTAL
            elif total_col_idx and mrp_col_idx < col_idx < total_col_idx:
                size_cols.append((clean_name, col_idx))

        # Sort sizes by column index
        size_cols.sort(key=lambda x: x[1])

        if not size_cols:
            raise DispatchMatrixParseError(
                "No size matrix columns detected. Expected size headers like 36..42 or S..XXL."
            )

        detected_sizes = [s[0] for s in size_cols]
        raw_lines: List[Dict[str, Any]] = []
        store_groups: Dict[str, List[Dict[str, Any]]] = {}
        total_data_rows = 0

        for r in range(header_row_idx + 1, ws.max_row + 1):
            store_val = ws.cell(r, store_col_idx).value
            if store_val is None or not str(store_val).strip():
                continue

            clean_store = str(store_val).strip()
            art_val = str(ws.cell(r, art_col_idx).value or "").strip()
            color_val = str(ws.cell(r, color_col_idx).value or "").strip()
            mrp_raw = ws.cell(r, mrp_col_idx).value

            if not art_val:
                continue

            total_data_rows += 1

            try:
                mrp = Decimal(str(mrp_raw or 0).strip())
            except Exception:
                mrp = Decimal("0.00")

            # Check reported total if available
            reported_total: Optional[int] = None
            if total_col_idx:
                tot_val = ws.cell(r, total_col_idx).value
                if tot_val is not None:
                    try:
                        reported_total = int(float(str(tot_val).strip()))
                    except Exception:
                        reported_total = None

            row_unpivoted_qty = 0

            # Unpivot sizes
            for size_name, s_col_idx in size_cols:
                q_val = ws.cell(r, s_col_idx).value
                if q_val is not None and str(q_val).strip() not in ("", "0", "None"):
                    try:
                        qty = int(float(str(q_val).strip()))
                    except (ValueError, TypeError):
                        qty = 0

                    if qty > 0:
                        row_unpivoted_qty += qty
                        item_code = f"{art_val}-{color_val}-{size_name}"
                        item_name = f"{art_val} {color_val} {size_name}"
                        
                        line_entry = {
                            "row_index": r,
                            "store_code": clean_store,
                            "article": art_val,
                            "color": color_val,
                            "size": size_name,
                            "quantity": qty,
                            "mrp": mrp,
                            "item_code": item_code,
                            "item_name": item_name,
                        }
                        raw_lines.append(line_entry)
                        
                        if clean_store not in store_groups:
                            store_groups[clean_store] = []
                        store_groups[clean_store].append(line_entry)

        return {
            "total_rows": total_data_rows,
            "detected_sizes": detected_sizes,
            "store_groups": store_groups,
            "raw_lines": raw_lines,
            "column_map": {
                "header_row": header_row_idx,
                "store_column": store_col_name,
                "article_column": art_col_name,
                "color_column": color_col_name,
                "mrp_column": mrp_col_name,
                "size_columns": detected_sizes,
            }
        }
