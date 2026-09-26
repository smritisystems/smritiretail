from collections import OrderedDict
from io import BytesIO
from typing import Any

try:
    from openpyxl import load_workbook
except Exception:  # pragma: no cover
    load_workbook = None


REQUIRED_HEADERS = {
    "article_no": "Article No (SAP)",
    "hsn_code": "HSN Code",
    "barcode": "EAN / Barcode",
    "style_article": "Vendor Style / Art",
    "material_description": "Material Description",
    "article_name": "Article Name",
    "color": "Color",
    "size": "Size",
    "quantity": "Quantity (Pairs)",
    "uom": "UOM",
    "mrp": "MRP (₹)",
    "base_rate": "Base Rate (₹)",
    "igst": "IGST (%)",
    "po_number": "PO Number",
}


def _text(value: Any) -> str:
    return str(value).strip() if value is not None else ""


def _number(value: Any, default: float = 0.0) -> float:
    try:
        return float(value) if value not in (None, "") else default
    except (TypeError, ValueError):
        return default


def _tax_percent(value: Any) -> float:
    """Convert spreadsheet fractions such as 0.05 into item-master percentage points such as 5.0."""
    rate = _number(value, 5.0)
    return rate * 100 if 0 < rate <= 1 else rate


def parse_po_item_master_excel(file_bytes: bytes) -> list[dict[str, Any]]:
    """Parse the consolidated PO item sheet into grouped ItemCreateRequest-compatible payloads."""
    if load_workbook is None:
        raise RuntimeError("openpyxl is required to read PO Excel files")

    workbook = load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)
    worksheet = workbook["All POs Items Master"] if "All POs Items Master" in workbook.sheetnames else workbook.active
    rows = worksheet.iter_rows(values_only=True)
    headers = [_text(value) for value in next(rows, ())]
    header_index = {header: index for index, header in enumerate(headers) if header}
    missing = [header for header in REQUIRED_HEADERS.values() if header not in header_index]
    if missing:
        raise ValueError(f"PO item sheet is missing required columns: {', '.join(missing)}")

    grouped: OrderedDict[str, dict[str, Any]] = OrderedDict()
    for row in rows:
        style_article = _text(row[header_index[REQUIRED_HEADERS["style_article"]]]).upper()
        if not style_article:
            continue

        article_no = _text(row[header_index[REQUIRED_HEADERS["article_no"]]]).upper()
        barcode = _text(row[header_index[REQUIRED_HEADERS["barcode"]]])
        color = _text(row[header_index[REQUIRED_HEADERS["color"]]])
        size = _text(row[header_index[REQUIRED_HEADERS["size"]]])
        variant_sku = article_no or f"{style_article}-{color}-{size}".strip("-").upper()
        variant_key = (variant_sku, barcode)
        item = grouped.setdefault(
            style_article,
            {
                "item_code": style_article,
                "item_name": _text(row[header_index[REQUIRED_HEADERS["article_name"]]])
                or _text(row[header_index[REQUIRED_HEADERS["material_description"]]])
                or style_article,
                "category": "Footwear",
                "style_code": style_article,
                "hsn_code": _text(row[header_index[REQUIRED_HEADERS["hsn_code"]]]) or "64041990",
                "tax_rate": _tax_percent(row[header_index[REQUIRED_HEADERS["igst"]]]),
                "primary_uom": _text(row[header_index[REQUIRED_HEADERS["uom"]]]) or "PAIR",
                "mrp": _number(row[header_index[REQUIRED_HEADERS["mrp"]]]),
                "selling_price": _number(row[header_index[REQUIRED_HEADERS["base_rate"]]]),
                "cost_price": _number(row[header_index[REQUIRED_HEADERS["base_rate"]]]),
                "attributes_json": {"source": "PURCHASE_ORDER_EXCEL", "po_numbers": []},
                "variants": [],
                "_variant_keys": set(),
            },
        )
        po_number = _text(row[header_index[REQUIRED_HEADERS["po_number"]]])
        if po_number and po_number not in item["attributes_json"]["po_numbers"]:
            item["attributes_json"]["po_numbers"].append(po_number)

        if variant_key in item["_variant_keys"]:
            continue
        item["_variant_keys"].add(variant_key)
        item["variants"].append(
            {
                "variant_sku": variant_sku,
                "variant_name": " / ".join(part for part in (style_article, color, size) if part) or variant_sku,
                "attributes_json": {
                    "style_article": style_article,
                    "color": color,
                    "size": size,
                    "quantity_pairs": _number(row[header_index[REQUIRED_HEADERS["quantity"]]]),
                },
                "mrp": _number(row[header_index[REQUIRED_HEADERS["mrp"]]], item["mrp"]),
                "selling_price": _number(row[header_index[REQUIRED_HEADERS["base_rate"]]], item["selling_price"]),
                "cost_price": _number(row[header_index[REQUIRED_HEADERS["base_rate"]]], item["cost_price"]),
                "barcodes": ([{"barcode": barcode, "barcode_type": "EAN13", "is_primary": True}] if barcode else []),
            }
        )

    payloads = []
    for item in grouped.values():
        item.pop("_variant_keys", None)
        payloads.append(item)
    return payloads
