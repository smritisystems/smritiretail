"""
Project      : SMRITI Retail OS
Author       : Jawahar Ramkripal Mallah
Designation  : Chief Systems Architect & Creator
Email        : support@smritibooks.com
Websites     : smritibooks.com | erpnbook.com | aitdl.com
Version      : 6.40.1
Created      : 2026-09-18
Copyright    : © SMRITIBooks.com. All Rights Reserved.
License      : Proprietary Commercial Software
Classification: B2B Dispatch & Tax Invoicing Studio Automated Test Battery
"""

import ast
import io
import pytest
import openpyxl
from decimal import Decimal
from datetime import date
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.services.dispatch_matrix_parser import DispatchMatrixParser, DispatchMatrixParseError
from app.services.dispatch_invoicing_engine import DispatchInvoicingEngine, DISPATCH_FROM_SNAPSHOT
from app.services.dispatch_artifact_pipeline import DispatchArtifactPipeline
from app.schemas.dispatch_invoicing import DispatchBatchGenerateRequest
from app.models.sales import SalesInvoice, SalesInvoiceItem
from app.models.distribution import EWayBill


def create_sample_dispatch_excel(sizes=None, stores=None) -> bytes:
    """Helper to construct an in-memory sample dispatch workbook."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Dispatch_Test"

    sizes = sizes or [36, 37, 38, 39, 40, 41, 42]
    stores = stores or ["TXAJ", "TW07"]

    # Header
    header = ["STORE NAME", "ARTICLE", "COLOR", "MRP"] + sizes + ["TOTAL"]
    ws.append(header)

    # Sample data rows
    for store in stores:
        row1 = [store, "CH-10-B", "MUSTARD", 2299] + [1] * len(sizes) + [len(sizes)]
        row2 = [store, "CH-24-G", "BLACK", 2499] + [2] * len(sizes) + [2 * len(sizes)]
        ws.append(row1)
        ws.append(row2)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.mark.asyncio
async def test_dispatch_matrix_parser_dynamic_sizes():
    """Unit Test 1: Assert dynamic detection of numeric shoe sizes and unpivoting."""
    excel_bytes = create_sample_dispatch_excel(sizes=[36, 37, 38, 39, 40, 41, 42])
    res = DispatchMatrixParser.parse_workbook(excel_bytes, "Dispatch_Test")

    assert res["sheet_name"] == "Dispatch_Test"
    assert res["detected_sizes"] == ["36", "37", "38", "39", "40", "41", "42"]
    assert len(res["store_groups"]) == 2
    assert "TXAJ" in res["store_groups"]
    assert "TW07" in res["store_groups"]

    # Check unpivoted items for TXAJ
    txaj_items = res["store_groups"]["TXAJ"]
    # 2 rows * 7 sizes = 14 unpivoted lines
    assert len(txaj_items) == 14
    first_item = txaj_items[0]
    assert first_item["store_code"] == "TXAJ"
    assert first_item["article"] == "CH-10-B"
    assert first_item["color"] == "MUSTARD"
    assert first_item["size"] == "36"
    assert first_item["quantity"] == 1
    assert first_item["mrp"] == Decimal("2299")
    assert first_item["item_code"] == "CH-10-B-MUSTARD-36"


@pytest.mark.asyncio
async def test_dispatch_matrix_parser_apparel_sizes():
    """Unit Test 2: Assert dynamic detection of apparel size matrix (S, M, L, XL)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Apparel"
    ws.append(["SITE CODE", "STYLE", "SHADE", "TAG PRICE", "S", "M", "L", "XL", "TOTAL"])
    ws.append(["TW97", "TOP-01", "BLUE", 1499, 2, 4, 4, 2, 12])

    buf = io.BytesIO()
    wb.save(buf)

    res = DispatchMatrixParser.parse_workbook(buf.getvalue(), "Apparel")
    assert res["detected_sizes"] == ["S", "M", "L", "XL"]
    assert len(res["store_groups"]["TW97"]) == 4
    total_qty = sum(it["quantity"] for it in res["store_groups"]["TW97"])
    assert total_qty == 12


@pytest.mark.asyncio
async def test_preflight_audit_calculation_and_tax(db_session: AsyncSession):
    """Unit Test 3: Validate Pre-Flight Dry-Run calculations and statutory GST arbitration."""
    excel_bytes = create_sample_dispatch_excel(sizes=[36, 37, 38], stores=["TXAJ", "TW07"])
    audit_res = await DispatchInvoicingEngine.run_preflight_audit(
        db=db_session,
        file_bytes=excel_bytes,
        sheet_name="Dispatch_Test",
        discount_pct=Decimal("43.76"),
        gst_rate=Decimal("5.00")
    )

    assert audit_res.total_stores == 2
    # Each store has 2 rows * 3 sizes = 6 lines
    # Row 1: 3 pairs, Row 2: 6 pairs = 9 pairs per store -> 18 pairs total
    assert audit_res.total_pairs == 18
    assert audit_res.is_valid_to_generate is True
    assert audit_res.audit_token.startswith("audit-")

    # TXAJ is Tamil Nadu (33) -> Interstate (depot is 27) -> IGST > 0, CGST == 0, SGST == 0
    txaj_store = next(s for s in audit_res.stores if s.store_code == "TXAJ")
    assert txaj_store.state_code == 33
    assert txaj_store.is_interstate is True
    assert txaj_store.igst_amount > 0
    assert txaj_store.cgst_amount == 0
    assert txaj_store.sgst_amount == 0
    assert txaj_store.grand_total > txaj_store.taxable_value


@pytest.mark.asyncio
async def test_execute_batch_invoicing_and_identity(db_session: AsyncSession):
    """Integration Test 4: Atomically allocate sovereign document numbers and persist in PostgreSQL."""
    excel_bytes = create_sample_dispatch_excel(sizes=[38, 39], stores=["TXAJ"])
    audit_res = await DispatchInvoicingEngine.run_preflight_audit(
        db=db_session,
        file_bytes=excel_bytes,
        sheet_name="Dispatch_Test",
        discount_pct=Decimal("43.76")
    )

    req = DispatchBatchGenerateRequest(
        audit_token=audit_res.audit_token,
        invoice_date="2026-09-05",
        series_prefix="TT2026-2027/",
        customer_name="Reliance Retail Limited",
        discount_pct=Decimal("43.76"),
        starting_sequence=990
    )

    records, raw_bytes = await DispatchInvoicingEngine.execute_batch(
        db=db_session,
        req=req,
        company_id="comp-001",
        branch_id="main",
        operator="test-runner"
    )

    assert len(records) == 1
    rec = records[0]
    assert rec["invoice_no"] == "TT2026-2027/990"
    assert rec["store_code"] == "TXAJ"
    assert rec["identity_code"].startswith("SAL-INV-")
    assert rec["eway_identity_code"].startswith("TAX-EWB-")

    # Verify directly in PostgreSQL database
    stmt = select(SalesInvoice).where(SalesInvoice.invoice_no == "TT2026-2027/990")
    res = await db_session.execute(stmt)
    inv = res.scalars().first()
    assert inv is not None
    assert inv.identity_code == rec["identity_code"]
    assert inv.customer_name == "Reliance Retail Limited"
    assert inv.place_of_supply_code == "33"

    # Verify items
    stmt_items = select(SalesInvoiceItem).where(SalesInvoiceItem.invoice_id == inv.id)
    res_items = await db_session.execute(stmt_items)
    items = res_items.scalars().all()
    assert len(items) == 4  # 2 rows * 2 sizes


@pytest.mark.asyncio
async def test_dispatch_artifact_pipeline_packaging(db_session: AsyncSession):
    """Integration Test 5: Verify Master Excel, Stamped Excel, and ZIP Bundle Packaging."""
    excel_bytes = create_sample_dispatch_excel(sizes=[38], stores=["TW07"])
    audit_res = await DispatchInvoicingEngine.run_preflight_audit(
        db=db_session,
        file_bytes=excel_bytes,
        sheet_name="Dispatch_Test"
    )

    req = DispatchBatchGenerateRequest(
        audit_token=audit_res.audit_token,
        invoice_date="2026-09-05",
        series_prefix="TT2026-2027/",
        starting_sequence=995
    )

    records, raw_bytes = await DispatchInvoicingEngine.execute_batch(
        db=db_session,
        req=req
    )

    # 1. Summary Excel
    excel_summary = DispatchArtifactPipeline.generate_excel_summary_matrix(records)
    wb_sum = openpyxl.load_workbook(io.BytesIO(excel_summary))
    assert "Invoice_Summary" in wb_sum.sheetnames
    assert "All_Items_Consolidated" in wb_sum.sheetnames

    # 2. Stamped Source Excel
    stamped_excel = DispatchArtifactPipeline.stamp_source_excel(raw_bytes, "Dispatch_Test", records)
    wb_stamp = openpyxl.load_workbook(io.BytesIO(stamped_excel))
    ws_stamp = wb_stamp["Dispatch_Test"]
    # Check Column M header
    assert ws_stamp.cell(1, 13).value == "Invoice details"
    assert ws_stamp.cell(2, 13).value == "Tax_Invoice_TT2026-2027_995"

    # 3. ZIP Packaging
    zip_bytes = await DispatchArtifactPipeline.build_delivery_zip(
        invoice_records=records,
        raw_file_bytes=raw_bytes,
        sheet_name="Dispatch_Test",
        batch_id="test-batch-001"
    )
    assert len(zip_bytes) > 0
    import zipfile
    with zipfile.ZipFile(io.BytesIO(zip_bytes)) as zf:
        namelist = zf.namelist()
        assert any(n.startswith("Tax_Invoice_PDFs/") for n in namelist)
        assert any(n.startswith("Eway_JSON/") for n in namelist)
        assert any(n.startswith("Tax_Invoice_Summary_") for n in namelist)
        assert any(n.startswith("Dispatch_Invoiced_") for n in namelist)


def test_rule13_dispatch_engine_identity_ast_compliance():
    """Governance Test 6: Statically inspect AST asserting zero ad-hoc ID generators in dispatch modules."""
    files_to_check = [
        "backend/app/services/dispatch_matrix_parser.py",
        "backend/app/services/dispatch_invoicing_engine.py",
        "backend/app/services/dispatch_artifact_pipeline.py",
        "backend/app/api/v1/dispatch_invoicing.py"
    ]

    for fpath in files_to_check:
        with open(fpath, "r", encoding="utf-8") as f:
            tree = ast.parse(f.read(), filename=fpath)

        for node in ast.walk(tree):
            if isinstance(node, ast.FunctionDef):
                assert node.name != "_uid", f"Rule 13 violation: Forbidden '_uid()' defined in {fpath}"
                assert node.name != "generate_uuid", f"Rule 13 violation: Forbidden generator in {fpath}"
